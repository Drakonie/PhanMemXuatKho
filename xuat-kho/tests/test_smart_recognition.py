"""Recognize real-world header variations without swallowing incomplete goods."""
import io
import sys
from pathlib import Path
import unittest
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from engine import analyze, export
from image_import import build_image_workbook
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]

def source(rows, merges=()):
    book = Workbook()
    for row in rows: book.active.append(row)
    for merged in merges: book.active.merge_cells(merged)
    stream = io.BytesIO(); book.save(stream)
    return stream.getvalue()

class SmartRecognitionTests(unittest.TestCase):
    def test_compact_unicode_headers_shuffled_and_dotted_date(self):
        result = analyze(source([
            ['Khách hàng: Bếp ăn A'], ['Ngày xuất: 02.10.2026'],
            ['ThànhTiền', 'Tên\u200bHàngHóa', 'ĐơnGiá', 'ＳｏＬｕｏｎｇ', 'ĐVT'],
            [20000, 'Gạo', 10000, 2, 'kg'],
        ]))[0]
        self.assertEqual(result['header'], 3)
        self.assertEqual(result['date'], '2026-10-02')
        self.assertEqual(result['mapping']['quantity'], 'D')
        self.assertEqual(result['total'], 20000)
        self.assertFalse(result['errors'])

    def test_three_header_rows_and_template_export(self):
        result = analyze(source([
            ['Khách hàng: Bếp ăn A'], ['Ngày xuất: 02/10/2026'],
            ['Tên', 'Số', 'Đơn', 'Thành'], ['hàng', 'lượng', 'giá', 'tiền'],
            ['hóa', 'thực xuất', 'VND', 'VND'], ['Gạo', 2, 10000, 20000],
        ]))[0]
        self.assertEqual((result['header'], result['header_span']), (3, 3))
        self.assertEqual(result['items'][0]['row'], 6)
        self.assertFalse(result['errors'])
        book = load_workbook(io.BytesIO(export((ROOT/'templates/Template.xlsx').read_bytes(), [result])))
        self.assertEqual(book.active['H27'].value, 20000)

    def test_nested_merged_actual_quantity(self):
        result = analyze(source([
            ['Khách hàng: Bếp ăn A'], ['Ngày xuất: 02/10/2026'],
            ['Tên hàng', 'Số lượng', None, 'Đơn giá', 'Thành tiền'],
            [None, 'Theo chứng từ', 'Thực', None, None],
            [None, None, 'xuất', 'VND', 'VND'],
            ['Gạo', 5, 2, 10000, 20000],
        ], ('A3:A5', 'B3:C3', 'B4:B5')))[0]
        self.assertEqual(result['header_span'], 3)
        self.assertEqual(result['mapping']['quantity'], 'C')
        self.assertEqual(result['total'], 20000)
        self.assertFalse(result['errors'])

    def test_partial_goods_after_header_cannot_be_hidden(self):
        result = analyze(source([
            ['Tên hàng', 'Số lượng', 'Đơn giá', 'Thành tiền'],
            ['Gạo', 'chưa rõ', 10000, 20000], ['Rau', 2, 10000, 20000],
        ]))[0]
        self.assertEqual(result['header_span'], 1)
        self.assertTrue(any('Dòng 2' in message for message in result['errors']))

    def test_three_rows_can_be_selected_manually(self):
        data = source([['Tên', 'Số', 'Đơn', 'Thành'], ['hàng', 'lượng', 'giá', 'tiền'],
                       ['hóa', 'thực xuất', 'VND', 'VND'], ['Gạo', 2, 10000, 20000]])
        result = analyze(data, {'Sheet': {'header': 1, 'header_span': 3,
            'mapping': {'name':'A', 'quantity':'B', 'price':'C', 'amount':'D'}}})[0]
        self.assertEqual(result['items'][0]['row'], 4)
        self.assertEqual(result['total'], 20000)

    def test_ocr_compact_headers_keep_shuffled_columns(self):
        data, report = build_image_workbook([{'filename':'image.png', 'text':
            'Khách hàng: Bếp ăn A\nNgày xuất: 02.10.2026\n'
            'ThanhTien | TenHangHoa | DonGia | SoLuong\n20000 | Gạo | 10000 | 2'}])
        result = analyze(data)[0]
        self.assertEqual(result['items'][0]['name'], 'Gạo')
        self.assertEqual(result['total'], 20000)
        self.assertFalse(result['errors'])

if __name__ == '__main__': unittest.main()
