import copy
import io
import json
import sys
import threading
import unittest
from pathlib import Path
from urllib.error import HTTPError
from urllib.request import Request, urlopen
from http.server import ThreadingHTTPServer
from zipfile import ZipFile
from openpyxl import Workbook, load_workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from app import Handler, SESSIONS, apply_settings, create_download
from engine import analyze
ROOT=Path(__file__).resolve().parents[1]
TEMPLATE=(ROOT/'templates/Template.xlsx').read_bytes()
SOURCE=(ROOT/'samples/data-mau.xlsx').read_bytes()

class QuietHandler(Handler):
    def log_message(self,*args): pass

class ApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.server=ThreadingHTTPServer(('127.0.0.1',0),QuietHandler)
        cls.thread=threading.Thread(target=cls.server.serve_forever,daemon=True);cls.thread.start()
        cls.url=f'http://127.0.0.1:{cls.server.server_port}'
    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown();cls.server.server_close();cls.thread.join()
        SESSIONS.clear()
    def request(self,path,payload=None):
        if payload is None: request=Request(self.url+path)
        else: request=Request(self.url+path,json.dumps(payload).encode(),{'Content-Type':'application/json'})
        try:
            with urlopen(request,timeout=10) as response:return response.status,response.read()
        except HTTPError as error:return error.code,error.read()
    def upload(self,source=SOURCE):
        boundary='tests-boundary-129'
        body=b'--'+boundary.encode()+b'\r\nContent-Disposition: form-data; name="data"; filename="test.xlsx"\r\nContent-Type: application/octet-stream\r\n\r\n'+source+b'\r\n--'+boundary.encode()+b'--\r\n'
        request=Request(self.url+'/api/analyze',body,{'Content-Type':'multipart/form-data; boundary='+boundary})
        with urlopen(request,timeout=10) as response:return json.loads(response.read())
    def test_health_upload_export_and_zip(self):
        status,body=self.request('/api/health');self.assertEqual(status,200);self.assertEqual(json.loads(body)['app'],'xuat-kho')
        data=self.upload();self.assertEqual(len(data['sheets']),4)
        status,body=self.request('/api/export',{'token':data['token'],'sheets':data['sheets']})
        self.assertEqual(status,200)
        wb=load_workbook(io.BytesIO(body));self.assertEqual(len(wb.worksheets),3)
        self.assertEqual(sum(s['H27'].value for s in wb),1569000)
        status,body=self.request('/api/export',{'token':data['token'],'sheets':data['sheets'],'format':'zip'})
        self.assertEqual(status,200)
        with ZipFile(io.BytesIO(body)) as archive:
            self.assertEqual(len(archive.namelist()),3)
            for name in archive.namelist():self.assertEqual(len(load_workbook(io.BytesIO(archive.read(name))).worksheets),1)
    def test_edits_are_validated_and_exported(self):
        data=self.upload();sheets=data['sheets'];sheets[0]['items'][0].update(quantity=5,amount=600000)
        status,body=self.request('/api/export',{'token':data['token'],'sheets':sheets})
        self.assertEqual(status,200);s=load_workbook(io.BytesIO(body)).active
        self.assertEqual(s['F12'].value,5);self.assertEqual(s['H27'].value,677000)
        sheets[0]['items'][0]['amount']=1
        status,body=self.request('/api/export',{'token':data['token'],'sheets':sheets});self.assertEqual(status,400)
    def test_cannot_silently_drop_rows_or_send_negative(self):
        data=self.upload();sheets=data['sheets'];sheets[0]['items'].pop()
        status,_=self.request('/api/export',{'token':data['token'],'sheets':sheets});self.assertEqual(status,400)
        data=self.upload();data['sheets'][0]['items'][0]['quantity']=-5
        status,_=self.request('/api/export',{'token':data['token'],'sheets':data['sheets']});self.assertEqual(status,400)
    def test_explicit_row_exclusion_is_allowed_and_unknown_ids_rejected(self):
        data=self.upload();data['sheets'][0]['excluded_rows']=[6]
        status,body=self.request('/api/export',{'token':data['token'],'sheets':data['sheets']})
        self.assertEqual(status,200)
        self.assertEqual(load_workbook(io.BytesIO(body)).active['H27'].value,77000)
        data['sheets'][0]['excluded_rows']=[999999]
        status,_=self.request('/api/export',{'token':data['token'],'sheets':data['sheets']});self.assertEqual(status,400)

    def test_invalid_session_and_remap(self):
        status,_=self.request('/api/export',{'token':'bad','sheets':[]});self.assertEqual(status,400)
        data=self.upload()
        status,body=self.request('/api/remap',{'token':data['token'],'overrides':{'Trung tâm':{'header':5,'mapping':{'name':'B','unit':'C','quantity':'D','price':'E','amount':'F','note':'G'}}}})
        self.assertEqual(status,200);self.assertEqual(json.loads(body)['sheets'][0]['total'],593000)
    def test_unparsed_source_error_remains_blocking(self):
        wb=Workbook();s=wb.active;s.append(['Khách hàng: ABC']);s.append(['Ngày 01 tháng 10 năm 2026']);s.append(['Tên hàng','Số lượng','Đơn giá','Thành tiền']);s.append(['A',2,10000,20000]);s.append(['B',None,10000,10000])
        stream=io.BytesIO();wb.save(stream)
        data=self.upload(stream.getvalue())
        status,body=self.request('/api/export',{'token':data['token'],'sheets':data['sheets']});self.assertEqual(status,400)
    def test_can_correct_recoverable_amount_mismatch(self):
        wb=Workbook();s=wb.active;s.append(['Khách hàng: ABC']);s.append(['Ngày 01 tháng 10 năm 2026']);s.append(['Tên hàng','Số lượng','Đơn giá','Thành tiền']);s.append(['A',2,10000,99999])
        stream=io.BytesIO();wb.save(stream)
        data=self.upload(stream.getvalue());self.assertTrue(data['sheets'][0]['errors'])
        data['sheets'][0]['items'][0]['amount']=20000
        status,body=self.request('/api/export',{'token':data['token'],'sheets':data['sheets']});self.assertEqual(status,200)
        self.assertEqual(load_workbook(io.BytesIO(body)).active['H27'].value,20000)

if __name__=='__main__':unittest.main()
