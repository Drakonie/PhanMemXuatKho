import copy
import io
import sys
import unittest
from pathlib import Path
from decimal import Decimal
from openpyxl import Workbook, load_workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from engine import analyze, export, number, money_words
ROOT=Path(__file__).resolve().parents[1]
TEMPLATE=(ROOT/'templates/Template.xlsx').read_bytes()

def bytes_of(wb):
    stream=io.BytesIO();wb.save(stream);return stream.getvalue()

def source(count=2):
    wb=Workbook();s=wb.active;s.title='Đơn hàng'
    s.append(['Khách hàng:', 'Công ty ABC']);s.append(['Ngày 01 tháng 10 năm 2026'])
    s.append(['Đơn giá (VNĐ)','Mặt hàng','Thành tiền','SL','Đơn vị tính'])
    for i in range(count): s.append([10000,'Hàng '+str(i+1),'=A%d*D%d'%(i+4,i+4),2,'kg'])
    s.append([None,'Tổng cộng']);s.append([None,'Người nhận hàng'])
    return wb

class EngineTests(unittest.TestCase):
    def test_sample_exact_totals_and_template(self):
        sheets=analyze((ROOT/'samples/data-mau.xlsx').read_bytes())
        self.assertEqual([s['total'] for s in sheets],[593000,432000,544000,1569000])
        self.assertFalse(sheets[-1]['selected'])
        self.assertTrue(all(not s['errors'] for s in sheets))
        wb=load_workbook(io.BytesIO(export(TEMPLATE,sheets)))
        self.assertEqual(len(wb.worksheets),3)
        s=wb.worksheets[0]
        self.assertEqual(s['B7'].value,sheets[0]['customer'])
        self.assertEqual(s['I14'].value,'Rau Nhà trường có')
        self.assertEqual(s['H27'].value,593000)
        self.assertEqual(s['B29'].value,'Năm trăm chín mươi ba nghìn đồng.')
        self.assertIsNone(s['B15'].value)
        self.assertEqual(s['A1'].value,load_workbook(io.BytesIO(TEMPLATE)).active['A1'].value)
        self.assertIn('B31:C32',[str(m) for m in s.merged_cells.ranges])
        self.assertEqual(s['E6'].value.strftime('%Y-%m-%d'),'2026-10-01')
    def test_shuffled_columns_uncached_formulas_and_footer(self):
        sheets=analyze(bytes_of(source()))
        self.assertEqual(sheets[0]['mapping']['name'],'B')
        self.assertEqual(sheets[0]['customer'],'Công ty ABC')
        self.assertEqual(sheets[0]['total'],40000)
        self.assertEqual(sheets[0]['errors'],[])
    def test_more_than_fifteen_rows_keeps_footer(self):
        sheets=analyze(bytes_of(source(22)))
        wb=load_workbook(io.BytesIO(export(TEMPLATE,sheets)))
        s=wb.active
        self.assertEqual(s['B33'].value,'Hàng 22')
        self.assertEqual(s['H34'].value,440000)
        self.assertEqual(s['B36'].value,'Bốn trăm bốn mươi nghìn đồng.')
        self.assertIn('B38:C39',[str(m) for m in s.merged_cells.ranges])
        self.assertEqual(s['B38'].value,'Người lập phiếu\n(Ký, họ tên)')
        self.assertEqual(copy.copy(s['B33'].border),copy.copy(s['B12'].border))
    def test_missing_and_conflicting_amount_block_export(self):
        wb=source();wb.active['D4']=None
        sheets=analyze(bytes_of(wb));self.assertTrue(sheets[0]['errors'])
        with self.assertRaises(ValueError):export(TEMPLATE,sheets)
        wb=source();wb.active['C4']=12345
        sheets=analyze(bytes_of(wb));self.assertTrue(sheets[0]['errors'])
        with self.assertRaises(ValueError):export(TEMPLATE,sheets)
    def test_derive_price_and_manual_mapping(self):
        wb=source();s=wb.active;s['A3']='Giá không chuẩn';s['C4']=20000;s['C5']=20000
        sheets=analyze(bytes_of(wb))
        self.assertEqual(sheets[0]['items'][0]['price'],10000)
        s['B3']='Hàng';sheets=analyze(bytes_of(wb));self.assertFalse(sheets[0]['selected'])
        sheets=analyze(bytes_of(wb),{'Đơn hàng':{'header':3,'mapping':{'name':'B','quantity':'D','price':'A','amount':'C','unit':'E'}}})
        self.assertEqual(sheets[0]['total'],40000)
    def test_vietnamese_numbers_money_words(self):
        self.assertEqual(number('120.000'),Decimal(120000))
        self.assertEqual(number('4,3'),Decimal('4.3'))
        self.assertEqual(number('1.234,50 đ'),Decimal('1234.50'))
        self.assertEqual(number('1,234.50'),Decimal('1234.50'))
        self.assertEqual(money_words(5051100),'Năm triệu không trăm năm mươi mốt nghìn một trăm đồng.')
        self.assertEqual(money_words(0),'Không đồng.')
        self.assertEqual(money_words(1000001),'Một triệu không trăm lẻ một đồng.')
    def test_zero_filter_and_formula_text(self):
        sheets=analyze((ROOT/'samples/data-mau.xlsx').read_bytes())
        sheets[0]['customer']='=1+1'
        wb=load_workbook(io.BytesIO(export(TEMPLATE,sheets,exclude_zero=True)))
        self.assertEqual(wb.active['B7'].data_type,'s')
        self.assertIsNone(wb.active['B14'].value)
        self.assertEqual(wb.active['H27'].value,593000)
    def test_cross_sheet_formula_without_cached_value(self):
        wb=source();s=wb.create_sheet('Tổng hợp')
        s.append(['Khách hàng: Công ty ABC']);s.append(['Ngày 01 tháng 10 năm 2026'])
        s.append(['Tên hàng','Số lượng','Đơn giá','Thành tiền'])
        s.append(['Hàng',"='Đơn hàng'!D4+'Đơn hàng'!D5",10000,'=B4*C4'])
        sheets=analyze(bytes_of(wb));self.assertEqual(sheets[1]['total'],40000)
        self.assertFalse(sheets[1]['selected'])

if __name__=='__main__': unittest.main()
