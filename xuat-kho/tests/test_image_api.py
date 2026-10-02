import base64
import io
import json
import sys
import threading
import unittest
from http.server import ThreadingHTTPServer
from pathlib import Path
from urllib.error import HTTPError
from urllib.request import Request, urlopen

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from api.index import handler as VercelHandler
from app import Handler, LOCK, SESSIONS, read_ocr_documents


def fixture_document(filename='phieu.png', missing_quantity=False):
    rows = [
        [(20, 'Khách hàng: Trường ABC')],
        [(20, 'Ngày 01 tháng 10 năm 2026')],
        [(20, 'Tên hàng'), (280, 'ĐVT'), (400, 'Số lượng'), (540, 'Đơn giá'), (700, 'Thành tiền')],
        [(20, 'Thịt heo'), (280, 'kg'), (400, '2'), (540, '100.000'), (700, '200.000')],
        [(20, 'Rau củ'), (280, 'kg'), (400, '1'), (540, '30.000'), (700, '30.000')],
    ]
    if missing_quantity:
        rows[4] = [cell for cell in rows[4] if cell[0] != 400]
    tsv = ['level\tpage_num\tblock_num\tpar_num\tline_num\tword_num\tleft\ttop\twidth\theight\tconf\ttext']
    for line_number, cells in enumerate(rows, 1):
        for word_number, (left, text) in enumerate(cells, 1):
            tsv.append(f'5\t1\t1\t1\t{line_number}\t{word_number}\t{left}\t{40 * line_number}\t{max(20, len(text) * 8)}\t22\t97\t{text}')
    return {'filename': filename, 'text': '\n'.join(' '.join(text for _, text in row) for row in rows),
            'tsv': '\n'.join(tsv), 'width': 900, 'height': 400, 'rotation': 0, 'language': 'vie+eng'}


class QuietLocalHandler(Handler):
    def log_message(self, *args):
        pass


class ImageApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.servers = {}
        for name, handler in [('local', QuietLocalHandler), ('vercel', VercelHandler)]:
            server = ThreadingHTTPServer(('127.0.0.1', 0), handler)
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            cls.servers[name] = (server, thread)

    @classmethod
    def tearDownClass(cls):
        for server, thread in cls.servers.values():
            server.shutdown()
            server.server_close()
            thread.join()
        with LOCK:
            SESSIONS.clear()

    def request(self, mode, path, documents=None, source=None, payload=None, overrides=None):
        server, _ = self.servers[mode]
        url = f'http://127.0.0.1:{server.server_port}{path}'
        if documents is None and source is None and payload is None:
            request = Request(url)
        elif mode == 'local' and payload is not None:
            request = Request(url, json.dumps(payload).encode(), {'Content-Type': 'application/json'})
        else:
            boundary = 'ocr-endpoint-test-129'
            fields = []
            if documents is not None:
                fields.append(('ocr_documents', json.dumps(documents, ensure_ascii=False).encode(), None))
            if overrides is not None:
                fields.append(('image_overrides', json.dumps(overrides).encode(), None))
            if source is not None:
                fields.append(('data', source, 'source.xlsx'))
            if payload is not None:
                fields.append(('payload', json.dumps(payload).encode(), None))
            body = b''
            for name, value, filename in fields:
                disposition = f'Content-Disposition: form-data; name="{name}"'
                if filename:
                    disposition += f'; filename="{filename}"'
                body += f'--{boundary}\r\n{disposition}\r\n\r\n'.encode() + value + b'\r\n'
            body += f'--{boundary}--\r\n'.encode()
            request = Request(url, body, {'Content-Type': f'multipart/form-data; boundary={boundary}'})
        try:
            with urlopen(request, timeout=15) as response:
                return response.status, response.read()
        except HTTPError as error:
            return error.code, error.read()

    def image_upload(self, mode='vercel', documents=None):
        status, body = self.request(mode, '/api/image-analyze', documents=documents or [fixture_document()])
        self.assertEqual(status, 200, body.decode(errors='replace'))
        return json.loads(body)

    def test_health_advertises_browser_ocr(self):
        for mode in self.servers:
            with self.subTest(mode=mode):
                status, body = self.request(mode, '/api/health')
                data = json.loads(body)
                self.assertEqual(status, 200)
                self.assertTrue(data['capabilities']['image_input'])
                self.assertEqual(data['ocr_engine'], 'tesseract-browser')

    def test_real_image_ocr_import_and_stateless_export(self):
        data = self.image_upload()
        source = base64.b64decode(data['source_base64'], validate=True)
        self.assertEqual(data['transport'], 'stateless')
        self.assertEqual(data['source_type'], 'image')
        self.assertEqual(len(data['sheets']), 1)
        sheet = data['sheets'][0]
        self.assertFalse(sheet['selected'])
        self.assertTrue(sheet['requires_review'])
        self.assertEqual(sheet['source_type'], 'image')
        self.assertEqual(sheet['confidence_kind'], 'table-structure')
        self.assertEqual(sheet['total'], 230000)
        self.assertEqual(len(sheet['items']), 2)
        self.assertTrue(sheet['image_info'])
        self.assertTrue(sheet['warnings'])
        sheet.update(selected=True, customer='Trường ABC', date='2026-10-01')
        # New invocation has only the generated workbook, no OCR cache/session.
        status, body = self.request('vercel', '/api/export', source=source, payload={'sheets': data['sheets']})
        self.assertEqual(status, 200, body.decode(errors='replace')[:500])
        self.assertEqual(load_workbook(io.BytesIO(body)).active['H27'].value, 230000)

    def test_local_image_session_export_and_remap_keep_ocr_metadata(self):
        data = self.image_upload('local')
        sheet = data['sheets'][0]
        status, body = self.request('local', '/api/remap', payload={
            'token': data['token'], 'overrides': {sheet['sheet']: {
                'header': sheet['header'], 'header_span': sheet['header_span'], 'mapping': sheet['mapping'],
            }},
        })
        self.assertEqual(status, 200, body.decode(errors='replace'))
        remapped = json.loads(body)['sheets']
        self.assertEqual(remapped[0]['source_type'], 'image')
        self.assertTrue(remapped[0]['requires_review'])
        remapped[0].update(selected=True, customer='Trường ABC', date='2026-10-01')
        status, body = self.request('local', '/api/export', payload={'token': data['token'], 'sheets': remapped})
        self.assertEqual(status, 200, body.decode(errors='replace')[:500])
        self.assertEqual(load_workbook(io.BytesIO(body)).active['H27'].value, 230000)

    def test_multiple_images_make_separate_review_sheets(self):
        data = self.image_upload(documents=[fixture_document('phieu-1.png'), fixture_document('phieu-2.jpg')])
        self.assertEqual(len(data['sheets']), 2)
        self.assertEqual(len({sheet['sheet'] for sheet in data['sheets']}), 2)
        self.assertTrue(all(not sheet['selected'] and sheet['requires_review'] for sheet in data['sheets']))
        self.assertEqual(sum(sheet['total'] for sheet in data['sheets']), 460000)

    def test_recipient_is_not_an_image_filename(self):
        for mode in self.servers:
            self.assertEqual(self.image_upload(mode=mode)['sheets'][0]['recipient'], '')
            document = {'filename': 'phieu.png', 'edited_text': True, 'text':
                        'Khách hàng: Trường ABC\nNgày xuất: 01/10/2026\nNgười nhận hàng: Nguyễn An\n'
                        'Tên hàng | Số lượng | Đơn giá | Thành tiền\nGạo | 2 | 10000 | 20000'}
            self.assertEqual(self.image_upload(mode=mode, documents=[document])['sheets'][0]['recipient'], 'Nguyễn An')

    def test_other_image_column_settings_survive_text_reparse(self):
        documents = [fixture_document('phieu-1.png'), fixture_document('phieu-2.png')]
        overrides = {'phieu-1': {'header': 3, 'header_span': 1,
                                'mapping': {'name': 'A', 'quantity': 'D', 'price': 'E', 'amount': '', 'unit': 'C', 'code': '', 'note': ''}}}
        for mode in self.servers:
            status, body = self.request(mode, '/api/image-analyze', documents=documents, overrides=overrides)
            self.assertEqual(status, 200, body.decode(errors='replace'))
            data = json.loads(body)
            self.assertEqual(data['overrides'], overrides)
            self.assertEqual(data['sheets'][0]['detection_method'], 'manual')
            self.assertNotIn('amount', data['sheets'][0]['mapping'])
            self.assertEqual(data['sheets'][0]['total'], 230000)
            self.assertEqual(data['sheets'][1]['detection_method'], 'headers')
            if mode == 'local':
                self.assertEqual(SESSIONS[data['token']]['overrides'], overrides)

    def test_invalid_image_column_settings_are_rejected(self):
        for overrides in [{'bad-sheet': {'header': 3, 'mapping': {'name': 'A'}}},
                          {'phieu': {'header': True, 'mapping': {'name': 'A'}}},
                          {'phieu': {'header': 3, 'mapping': {'name': 'A;import'}}}]:
            status, _ = self.request('vercel', '/api/image-analyze', documents=[fixture_document()], overrides=overrides)
            self.assertEqual(status, 400)

    def test_missing_numeric_data_is_not_silently_exported(self):
        data = self.image_upload(documents=[fixture_document(missing_quantity=True)])
        self.assertTrue(data['sheets'][0]['errors'])
        data['sheets'][0].update(selected=True, customer='Trường ABC', date='2026-10-01')
        source = base64.b64decode(data['source_base64'])
        status, _ = self.request('vercel', '/api/export', source=source, payload={'sheets': data['sheets']})
        self.assertEqual(status, 400)

    def test_invalid_or_oversized_ocr_documents_are_rejected(self):
        invalid = [[], [{}], 'invalid', [fixture_document()] * 11]
        bad_dimensions = fixture_document()
        bad_dimensions['width'] = float('inf')
        invalid.append([bad_dimensions])
        for documents in invalid:
            with self.subTest(documents_type=type(documents).__name__):
                status, _ = self.request('vercel', '/api/image-analyze', documents=documents)
                self.assertEqual(status, 400)
        document = fixture_document()
        document['text'] = 'x' * 1_000_001
        with self.assertRaises(ValueError):
            read_ocr_documents(json.dumps([document]).encode())

    def test_ocr_asset_path_cannot_escape_asset_directory(self):
        status, _ = self.request('local', '/ocr/../../app.py')
        self.assertEqual(status, 404)


if __name__ == '__main__':
    unittest.main()
