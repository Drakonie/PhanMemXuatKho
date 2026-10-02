"""HTTP tests for remap/export across independent serverless invocations."""
import copy
import http.client
import io
import json
import sys
import threading
import unittest
from http.server import ThreadingHTTPServer
from pathlib import Path
from unittest.mock import patch
from urllib.error import HTTPError
from urllib.request import Request, urlopen
from zipfile import ZipFile

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from api.index import MAX_UPLOAD, UploadTooLarge, handler, read_request_body
from app import SESSIONS

ROOT = Path(__file__).resolve().parents[1]
SOURCE = (ROOT / 'samples/data-mau.xlsx').read_bytes()
TEMPLATE = (ROOT / 'templates/Template.xlsx').read_bytes()


class VercelTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.server = ThreadingHTTPServer(('127.0.0.1', 0), handler)
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()
        cls.url = f'http://127.0.0.1:{cls.server.server_port}'

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.server.server_close()
        cls.thread.join()

    def request(self, path, payload=None, source=SOURCE, template=None, raw=None, content_type=None):
        if raw is not None:
            request = Request(self.url + path, raw, {'Content-Type': content_type or 'application/json'})
        elif path == '/api/health' or path == '/template':
            request = Request(self.url + path)
        else:
            boundary = 'serverless-tests-129'
            fields = [('data', source, 'source.xlsx')]
            if template is not None:
                fields.append(('template', template, 'template.xlsx'))
            if payload is not None:
                fields.append(('payload', json.dumps(payload).encode(), None))
            body = b''
            for name, value, filename in fields:
                disposition = f'Content-Disposition: form-data; name="{name}"'
                if filename:
                    disposition += f'; filename="{filename}"'
                body += f'--{boundary}\r\n{disposition}\r\n\r\n'.encode() + value + b'\r\n'
            body += f'--{boundary}--\r\n'.encode()
            request = Request(self.url + path, body, {'Content-Type': f'multipart/form-data; boundary={boundary}'})
        try:
            with urlopen(request, timeout=15) as response:
                return response.status, response.read(), response.headers
        except HTTPError as error:
            return error.code, error.read(), error.headers

    def upload(self):
        status, body, _ = self.request('/api/analyze')
        self.assertEqual(status, 200, body.decode(errors='replace'))
        return json.loads(body)

    def test_health_and_template(self):
        status, body, _ = self.request('/api/health')
        data = json.loads(body)
        self.assertEqual(status, 200)
        self.assertEqual(data['transport'], 'stateless')
        self.assertEqual(data['max_upload_bytes'], MAX_UPLOAD)
        status, body, _ = self.request('/template')
        self.assertEqual(status, 200)
        self.assertEqual(body, TEMPLATE)

    def test_export_does_not_use_process_sessions(self):
        data = self.upload()
        self.assertEqual(data['transport'], 'stateless')
        # The token need not exist anywhere, even after a cold start.
        with patch.dict(SESSIONS, {}, clear=True):
            status, body, headers = self.request('/api/export', {
                'token': 'unrelated-process', 'sheets': data['sheets'], 'overrides': {},
            })
            self.assertEqual(SESSIONS, {})
        self.assertEqual(status, 200, body.decode(errors='replace')[:500])
        self.assertIn('Phieu-xuat-kho.xlsx', headers['Content-Disposition'])
        workbook = load_workbook(io.BytesIO(body))
        self.assertEqual(len(workbook.worksheets), 3)
        self.assertEqual(sum(sheet['H27'].value for sheet in workbook), 1569000)

    def test_custom_template_and_zip(self):
        data = self.upload()
        workbook = load_workbook(io.BytesIO(TEMPLATE))
        workbook.active['A4'] = 'Template tùy chỉnh'
        output = io.BytesIO()
        workbook.save(output)
        status, body, _ = self.request('/api/export', {'sheets': data['sheets'], 'format': 'zip'}, template=output.getvalue())
        self.assertEqual(status, 200)
        with ZipFile(io.BytesIO(body)) as archive:
            self.assertEqual(len(archive.namelist()), 3)
            for name in archive.namelist():
                workbook = load_workbook(io.BytesIO(archive.read(name)))
                self.assertEqual(workbook.active['A4'].value, 'Template tùy chỉnh')

    def test_cumulative_remaps_survive_following_export(self):
        data = self.upload()
        overrides = {}
        for sheet in data['sheets'][:2]:
            overrides[sheet['sheet']] = {'header': sheet['header'], 'header_span': sheet['header_span'], 'mapping': sheet['mapping']}
            status, body, _ = self.request('/api/remap', {'overrides': overrides})
            self.assertEqual(status, 200, body.decode(errors='replace'))
            result = json.loads(body)
        for name in overrides:
            self.assertEqual(next(sheet for sheet in result['sheets'] if sheet['sheet'] == name)['detection_method'], 'manual')
        status, body, _ = self.request('/api/export', {'overrides': overrides, 'sheets': result['sheets']})
        self.assertEqual(status, 200, body.decode(errors='replace')[:500])
        self.assertEqual(sum(sheet['H27'].value for sheet in load_workbook(io.BytesIO(body))), 1569000)

    def test_edits_keep_source_row_validation(self):
        data = self.upload()
        edited = copy.deepcopy(data['sheets'])
        edited[0]['items'][0].update(quantity=5, amount=600000)
        status, body, _ = self.request('/api/export', {'sheets': edited})
        self.assertEqual(status, 200)
        self.assertEqual(load_workbook(io.BytesIO(body)).active['H27'].value, 677000)
        edited[0]['items'].pop()
        status, _, _ = self.request('/api/export', {'sheets': edited})
        self.assertEqual(status, 400)

    def test_invalid_source_payload_mapping_and_private_path(self):
        status, _, _ = self.request('/api/analyze', source=b'not-an-xlsx')
        self.assertEqual(status, 400)
        status, _, _ = self.request('/api/export', {'sheets': []}, raw=b'{"token":"old"}')
        self.assertEqual(status, 400)
        status, _, _ = self.request('/api/remap', {'overrides': {'Trung tâm': {'header': True, 'mapping': {}}}})
        self.assertEqual(status, 400)
        status, _, _ = self.request('/samples/data-mau.xlsx', {})
        self.assertEqual(status, 404)

    def test_upload_and_download_limits_return_clear_errors(self):
        with patch('api.index.MAX_UPLOAD', 128):
            status, body, _ = self.request('/api/analyze', raw=b'x' * 129)
        self.assertEqual(status, 413)
        self.assertIn('4 MB', json.loads(body)['error'])
        data = self.upload()
        with patch('api.index.create_download', return_value=(b'x' * (MAX_UPLOAD + 1), 'application/zip', 'too-large.zip')):
            status, body, headers = self.request('/api/export', {'sheets': data['sheets']})
        self.assertEqual(status, 413)
        self.assertEqual(headers.get_content_type(), 'application/json')
        self.assertIsNone(headers.get('Content-Disposition'))
        self.assertIn('4 MB', json.loads(body)['error'])

    def test_actual_chunked_http_upload(self):
        boundary = 'chunked-browser-419'
        body = (
            f'--{boundary}\r\nContent-Disposition: form-data; name="data"; filename="sample.xlsx"\r\n\r\n'.encode()
            + SOURCE + f'\r\n--{boundary}--\r\n'.encode()
        )
        connection = http.client.HTTPConnection('127.0.0.1', self.server.server_port, timeout=15)
        try:
            # A generator forces chunked framing and omits Content-Length,
            # reproducing the production Fluid Compute proxy behavior.
            connection.request('POST', '/api/analyze', body=(body[start:start+2048] for start in range(0, len(body), 2048)),
                               headers={'Content-Type': f'multipart/form-data; boundary={boundary}'}, encode_chunked=True)
            response = connection.getresponse()
            result = response.read()
            self.assertEqual(response.status, 200, result.decode(errors='replace'))
            self.assertEqual(len(json.loads(result)['sheets']), 4)
        finally:
            connection.close()

    def test_buffered_body_without_length_and_chunk_extensions(self):
        self.assertEqual(read_request_body(io.BytesIO(b'body'), {}), b'body')
        body = b'3;version=1\r\nabc\r\n2\r\nde\r\n0\r\nExample: test\r\n\r\n'
        self.assertEqual(read_request_body(io.BytesIO(body), {'Transfer-Encoding': 'chunked'}), b'abcde')

    def test_chunked_limits_and_invalid_framing(self):
        with patch('api.index.MAX_UPLOAD', 4):
            with self.assertRaises(UploadTooLarge):
                read_request_body(io.BytesIO(b'3\r\nabc\r\n2\r\nde\r\n0\r\n\r\n'), {'Transfer-Encoding': 'chunked'})
            with self.assertRaises(UploadTooLarge):
                read_request_body(io.BytesIO(b'12345'), {})
        malformed = [b'nothex\r\n', b'2\r\na', b'2\r\nabXX0\r\n\r\n', b'0\r\n']
        for body in malformed:
            with self.subTest(body=body), self.assertRaises(ValueError):
                read_request_body(io.BytesIO(body), {'Transfer-Encoding': 'chunked'})
        with self.assertRaises(ValueError):
            read_request_body(io.BytesIO(b'0\r\n\r\n'), {'Content-Length': '0', 'Transfer-Encoding': 'chunked'})


if __name__ == '__main__':
    unittest.main()
