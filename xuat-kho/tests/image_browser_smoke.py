"""Real Vietnamese OCR, review, correction and template-export browser check.

Requires Playwright and Chromium. Run against a started local server with
``python tests/image_browser_smoke.py --base http://127.0.0.1:8097``.
The checked-in photo is clearly labelled synthetic and contains no customer data.
"""
from __future__ import annotations

import argparse
import os
from pathlib import Path
import tempfile
from zipfile import ZipFile

from openpyxl import load_workbook
from playwright.sync_api import expect, sync_playwright


ROOT = Path(__file__).resolve().parents[1]
FIXTURE = ROOT / 'samples/anh-minh-hoa.png'
EXPECTED = [
    ('Gạo tẻ', 'kg', 2, 20000, 40000),
    ('Thịt lợn', 'kg', 3, 80000, 240000),
    ('Rau cải', 'kg', 5, 12000, 60000),
]
TEXT = '''Khách hàng: Bếp ăn Minh An
Ngày xuất: 01/10/2026
Địa chỉ: 12 Đường Hoa, Hà Nội
Tên hàng | ĐVT | Số lượng | Đơn giá | Thành tiền
Gạo tẻ | kg | 3 | 20.000 | 60.000
Thịt lợn | kg | 3 | 80.000 | 240.000
Rau cải | kg | 5 | 12.000 | 60.000
Trứng gà | quả | 2 | 5.000 | 10.000
'''


def response_is_image(response):
    return response.url.split('?', 1)[0].endswith('/api/image-analyze')


def analyze(page):
    with page.expect_response(response_is_image, timeout=240000) as pending:
        page.locator('#analyze').click()
    response = pending.value
    assert response.status == 200, response.text()
    data = response.json()
    page.wait_for_function("document.getElementById('workspace').hidden === false && !document.getElementById('status').classList.contains('busy')")
    # Original photos stay in the browser. The request contains OCR words only.
    assert 'name="ocr_documents"' in response.request.post_data
    assert 'Content-Type: image/png' not in response.request.post_data
    assert data['source_type'] == 'image' and data['source_base64']
    return data


def reparse(page, text):
    page.locator('[data-tab=source]').click()
    page.locator('#ocr-text-draft').fill(text)
    with page.expect_response(response_is_image) as pending:
        page.locator('[data-reparse-ocr]').click()
    response = pending.value
    assert response.status == 200, response.text()
    page.wait_for_function("document.getElementById('status').textContent.includes('Đã đọc lại văn bản')")
    assert page.locator('.selected:checked').count() == 0
    assert not page.locator('[data-image-confirm]').is_checked()
    assert page.locator('#export').is_disabled()
    return response.json()


def download(page, destination):
    with page.expect_download() as pending:
        page.locator('#export').click()
    pending.value.save_as(destination)
    return destination


def assert_workbook(path, expected, total):
    workbook = load_workbook(path)
    assert len(workbook.worksheets) == 1
    sheet = workbook.active
    assert sheet['B7'].value == 'Bếp ăn Minh An'
    assert sheet['E6'].value.strftime('%Y-%m-%d') == '2026-10-01'
    assert sheet['H27'].value == total
    rows = [tuple(sheet.cell(index + 12, column).value for column in (2, 4, 5, 7, 8)) for index in range(len(expected))]
    assert rows == expected, rows


def run(base='http://127.0.0.1:8097'):
    with tempfile.TemporaryDirectory() as directory, sync_playwright() as playwright:
        temporary = Path(directory)
        launch = dict(headless=True, executable_path='/usr/bin/chromium', args=['--no-sandbox'])
        if base.startswith('https://') and os.environ.get('HTTPS_PROXY'):
            launch['proxy'] = {'server': os.environ['HTTPS_PROXY']}
        browser = playwright.chromium.launch(**launch)
        page = browser.new_page(viewport={'width': 1440, 'height': 1100}, locale='vi-VN')
        page_errors = []
        page.on('pageerror', lambda error: page_errors.append(str(error)))
        page.goto(base)
        page.locator('#mode-images').click()
        page.locator('#image-files').set_input_files(str(FIXTURE))
        assert page.locator('.image-thumbnail').count() == 1
        page.locator('[data-preview-image="0"]').click()
        expect(page.locator('#image-zoom')).to_be_visible()
        page.locator('#image-zoom-close').click()
        data = analyze(page)
        assert len(data['sheets']) == 1
        source = data['sheets'][0]
        actual = [(item['name'], item['unit'], item['quantity'], item['price'], item['amount']) for item in source['items']]
        assert actual == EXPECTED, actual
        assert source['customer'] == 'Bếp ăn Minh An' and source['date'] == '2026-10-01'
        assert source['recipient'] == '', 'An image filename must not become the recipient'
        assert not source['errors'], source['errors']
        assert source['selected'] is False and source['requires_review'] is True
        assert page.locator('.selected:checked').count() == 0
        assert page.locator('#export').is_disabled()
        page.locator('#select-details').click()
        assert page.locator('.selected:checked').count() == 0, 'Bulk selection must leave image results unconfirmed'
        page.locator('[data-image-confirm]').check()
        assert page.locator('.selected:checked').count() == 1
        assert_workbook(download(page, temporary / 'ocr-original.xlsx'), EXPECTED, 340000)
        page.screenshot(path=str(ROOT / 'samples/giao-dien-anh-v4.png'), full_page=True)
        page.set_viewport_size({'width': 390, 'height': 844})
        assert page.evaluate('document.documentElement.scrollWidth <= innerWidth'), 'Items overflow mobile viewport'
        page.locator('[data-tab=source]').click()
        structured = page.locator('#ocr-text-draft').input_value()
        assert 'Tên hàng | Mã hàng | ĐVT | Số lượng | Đơn giá | Thành tiền | Ghi chú' in structured
        assert 'Gạo tẻ |' in structured and '| 20.000 | 40.000 |' in structured
        assert page.evaluate('document.documentElement.scrollWidth <= innerWidth'), 'OCR text/TSV overflow mobile viewport'
        page.screenshot(path=str(ROOT / 'samples/giao-dien-anh-mobile-v4.png'), full_page=True)
        page.set_viewport_size({'width': 1440, 'height': 1100})
        unchanged = reparse(page, structured)
        assert unchanged['sheets'][0]['total'] == 340000 and not unchanged['sheets'][0]['errors']
        assert unchanged['sheets'][0]['ocr_confidence'] is None

        # Missing/uncertain quantity must remain an error, rather than exporting
        # the other three goods and silently losing the fourth row.
        uncertain = reparse(page, TEXT.replace('Trứng gà | quả | 2 |', 'Trứng gà | quả | ? |'))
        assert uncertain['sheets'][0]['errors']
        assert len(uncertain['sheets'][0]['items']) == 3
        page.locator('[data-image-confirm]').check()
        with page.expect_response(lambda response: response.url.endswith('/api/export')) as pending:
            page.locator('#export').click()
        assert pending.value.status == 400, 'Unknown OCR quantity must block export'
        expect(page.locator('#status')).to_have_class('status error')
        corrected = reparse(page, TEXT)
        assert not corrected['sheets'][0]['errors'], corrected['sheets'][0]['errors']
        assert len(corrected['sheets'][0]['items']) == 4
        page.locator('[data-image-confirm]').check()
        corrected_expected = [('Gạo tẻ', 'kg', 3, 20000, 60000), *EXPECTED[1:], ('Trứng gà', 'quả', 2, 5000, 10000)]
        assert_workbook(download(page, temporary / 'ocr-corrected.xlsx'), corrected_expected, 370000)
        page.locator('[data-tab=mapping]').click()
        page.locator('[data-action=remap]').click()
        page.wait_for_function("document.getElementById('status').textContent.includes('Đã đọc lại sheet')")
        assert page.locator('.selected:checked').count() == 0
        assert not page.locator('[data-image-confirm]').is_checked()
        assert page.locator('.item-table tbody tr').count() == 4

        # A second input image gets an independent review checkbox and receipt.
        second = temporary / 'anh-minh-hoa-2.png'
        second.write_bytes(FIXTURE.read_bytes())
        page.locator('#clear-images').click()
        page.locator('#image-files').set_input_files([str(FIXTURE), str(second)])
        data = analyze(page)
        assert len(data['sheets']) == 2
        assert all(len(sheet['items']) == 3 and not sheet['errors'] for sheet in data['sheets'])
        assert page.locator('.selected:checked').count() == 0
        page.locator('[data-image-confirm]').check()
        page.locator('[data-open-sheet="1"]').click()
        assert not page.locator('[data-image-confirm]').is_checked()
        page.locator('[data-image-confirm]').check()
        assert '680.000' in page.locator('#total').inner_text()
        # Correcting the second photo preserves a prior column choice and review
        # of the first photo; its amount is deliberately inferred from qty×price.
        page.locator('[data-open-sheet="0"]').click()
        first_name = page.evaluate('sheets[0].sheet')
        page.locator('[data-tab=mapping]').click()
        page.locator('[data-map=amount]').select_option('')
        page.locator('[data-action=remap]').click()
        page.wait_for_function("document.getElementById('status').textContent.includes('Đã đọc lại sheet')")
        page.locator('[data-image-confirm]').check()
        page.locator('[data-open-sheet="1"]').click()
        page.locator('[data-tab=source]').click()
        second_text = page.locator('#ocr-text-draft').input_value()
        with page.expect_response(response_is_image) as pending:
            page.locator('[data-reparse-ocr]').click()
        assert pending.value.status == 200, pending.value.text()
        page.wait_for_function("document.getElementById('status').textContent.includes('Đã đọc lại văn bản')")
        assert page.locator('.selected:checked').count() == 1
        assert not page.locator('[data-image-confirm]').is_checked()
        assert page.evaluate('(name)=>overrides[name].mapping.amount', first_name) == ''
        assert page.evaluate('sheets[0].detection_method') == 'manual'
        page.locator('[data-image-confirm]').check()
        assert '680.000' in page.locator('#total').inner_text()
        page.locator('#export-format').select_option('zip')
        with ZipFile(download(page, temporary / 'ocr-two-images.zip')) as archive:
            assert len(archive.namelist()) == 2
        assert not page_errors, page_errors
        browser.close()
        print('Image browser passed: actual Vietnamese OCR, local assets, photo preview, required per-image review, template export, mobile source view, uncertain-row blocking, editable-text recovery, remap confirmation reset, two images and ZIP.')


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--base', default='http://127.0.0.1:8097')
    run(parser.parse_args().base)
