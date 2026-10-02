"""Export recovery through the real UI and API, with stalled network fixtures.

The clock advances deadlines without making the test wait 90 seconds. Fixtures
ignore AbortSignal deliberately: UI recovery must also work if a body promise
never settles. Successful retries and manual downloads use real Excel/ZIP files.
"""
from __future__ import annotations

import argparse
from pathlib import Path
import tempfile
from zipfile import ZipFile

from openpyxl import load_workbook
from playwright.sync_api import expect, sync_playwright


ROOT = Path(__file__).resolve().parents[1]
OCR_TEXT = """Khách hàng: Bếp ăn Minh An
Ngày xuất: 01/10/2026
Tên hàng | ĐVT | Số lượng | Đơn giá | Thành tiền
Gạo tẻ | kg | 2 | 20.000 | 40.000
Thịt lợn | kg | 3 | 80.000 | 240.000
Rau cải | kg | 5 | 12.000 | 60.000
"""


def idle(page):
    page.wait_for_function("busy === ''")


def download(page, selector, destination):
    with page.expect_download() as pending:
        page.locator(selector).click()
    pending.value.save_as(destination)
    idle(page)
    return destination


def set_fixture(page, stage):
    page.evaluate("""stage => {
      window.__realExportFetch ??= window.fetch;
      window.__exportStage = stage;
      window.__exportSeen = '';
      window.fetch = (path, options) => {
        if (path !== '/api/export' || !window.__exportStage) return window.__realExportFetch(path, options);
        window.__exportSignal = options.signal;
        window.__exportSeen = window.__exportStage;
        if (window.__exportStage === 'headers') return new Promise(() => {});
        if (window.__exportStage === 'body') return Promise.resolve({
          ok: true, status: 200, headers: new Headers(), blob: () => new Promise(() => {})
        });
        if (window.__exportStage === 'invalid') return Promise.resolve(new Response('not an Excel file'));
        return window.__realExportFetch(path, options);
      };
    }""", stage)


def assert_recovered(page, snapshot):
    idle(page)
    expect(page.locator('#export')).to_be_enabled()
    expect(page.locator('[data-field=customer]')).to_be_enabled()
    expect(page.locator('#export-cancel')).to_be_hidden()
    assert page.evaluate('JSON.stringify(exportSettings())') == snapshot
    assert page.evaluate('window.__exportSignal.aborted')


def run(base):
    with tempfile.TemporaryDirectory() as directory, sync_playwright() as playwright:
        temporary = Path(directory)
        browser = playwright.chromium.launch(
            headless=True, executable_path='/usr/bin/chromium', args=['--no-sandbox'])
        page = browser.new_page(viewport={'width': 1440, 'height': 1050}, locale='vi-VN')
        errors = []
        page.on('pageerror', lambda error: errors.append(str(error)))
        page.goto(base)
        page.locator('#theme-normal').click()
        assert page.evaluate("getComputedStyle(document.body).backgroundColor") == 'rgb(255, 255, 255)'
        assert page.evaluate("getComputedStyle(document.body).backgroundImage") == 'none'
        page.locator('#data-file').set_input_files(str(ROOT / 'samples/data-mau.xlsx'))
        page.locator('#analyze').click()
        page.locator('#workspace').wait_for(state='visible')
        idle(page)
        page.locator('#deselect-all').click()
        # A label provides the larger hit area used by less experienced users.
        page.locator('.sheet-select-target').first.click()
        page.locator('[data-item="0"][data-item-field="note"]').fill('Giữ dữ liệu sau khi hủy hoặc hết thời gian')
        snapshot = page.evaluate('JSON.stringify(exportSettings())')
        total = page.evaluate('sheetTotal(current())')
        workbook = load_workbook(download(page, '#export', temporary / 'initial.xlsx'))
        assert len(workbook.worksheets) == 1
        assert workbook.active['H27'].value == total
        assert workbook.active['I12'].value == 'Giữ dữ liệu sau khi hủy hoặc hết thời gian'
        expect(page.locator('#export-download')).to_be_visible()
        previous_href = page.locator('#export-download').get_attribute('href')

        page.clock.install()
        # The manual file remains usable after the old 30-second expiry.
        page.clock.fast_forward(31000)
        manual = load_workbook(download(page, '#export-download', temporary / 'manual.xlsx'))
        assert manual.active['H27'].value == total

        for stage in ('headers', 'body'):
            set_fixture(page, stage)
            page.locator('#export').click()
            page.wait_for_function('window.__exportSeen === window.__exportStage')
            expect(page.locator('#export-cancel')).to_be_enabled()
            expect(page.locator('[data-field=customer]')).to_be_disabled()
            page.locator('#export-cancel').click()
            assert_recovered(page, snapshot)
            expect(page.locator('#status')).to_contain_text('Đã hủy tạo file')
            assert page.locator('#export-download').get_attribute('href') == previous_href

        for stage in ('headers', 'body'):
            set_fixture(page, stage)
            page.locator('#export').click()
            page.wait_for_function('window.__exportSeen === window.__exportStage')
            page.clock.fast_forward(90001)
            assert_recovered(page, snapshot)
            expect(page.locator('#status')).to_contain_text('90 giây')
            expect(page.locator('#status')).to_have_class('status error')
            assert page.locator('#export-download').get_attribute('href') == previous_href

        set_fixture(page, 'invalid')
        page.locator('#export').click()
        idle(page)
        expect(page.locator('#status')).to_contain_text('chưa phải Excel hoặc ZIP hợp lệ')
        assert page.locator('#export-download').get_attribute('href') == previous_href

        # The next real export succeeds without re-importing or losing edits.
        set_fixture(page, '')
        page.locator('#theme-linh').click()
        page.locator('#export-format').select_option('zip')
        with ZipFile(download(page, '#export', temporary / 'retry.zip')) as archive:
            assert len(archive.namelist()) == 1
            assert archive.testzip() is None
        assert page.locator('#export-download').get_attribute('href') != previous_href
        assert page.locator('#export-download').get_attribute('download').endswith('.zip')

        # If the browser blocks a programmatic click, the visible link works.
        page.evaluate("""() => {
          window.__anchorClick = HTMLAnchorElement.prototype.click;
          HTMLAnchorElement.prototype.click = function() {};
        }""")
        page.locator('#export-format').select_option('xlsx')
        page.locator('#export').click()
        idle(page)
        expect(page.locator('#export-download')).to_be_visible()
        workbook = load_workbook(download(page, '#export-download', temporary / 'blocked-click.xlsx'))
        assert workbook.active['H27'].value == total
        page.evaluate('() => {HTMLAnchorElement.prototype.click = window.__anchorClick;}')

        # An image result keeps its explicit OCR review and customer selection.
        page.evaluate("""text => {
          window.Tesseract = {createWorker: async () => ({
            setParameters: async () => {},
            recognize: async () => ({data: {text, confidence: 99}}),
            terminate: async () => {}
          })};
          ImageRegions.readTable = async () => null;
        }""", OCR_TEXT)
        page.locator('#mode-images').click()
        page.locator('#image-files').set_input_files(str(ROOT / 'samples/anh-minh-hoa.png'))
        page.locator('#analyze').click()
        idle(page)
        expect(page.locator('#export-download')).to_be_hidden()
        expect(page.locator('#export')).to_be_disabled()
        page.locator('[data-customer-select="0"]').check()
        snapshot = page.evaluate('JSON.stringify(exportSettings())')
        set_fixture(page, 'body')
        page.locator('#export').click()
        page.wait_for_function("window.__exportSeen === 'body'")
        page.locator('#export-cancel').click()
        assert_recovered(page, snapshot)
        assert page.evaluate('sheets[0]._ocrConfirmed && sheets[0].selected')
        set_fixture(page, '')
        page.locator('#theme-normal').click()
        workbook = load_workbook(download(page, '#export', temporary / 'image-retry.xlsx'))
        assert workbook.active['H27'].value == 340000
        assert not errors, errors
        browser.close()
    print('Export browser passed: Excel/ZIP, white Normal/Linh switch, header/body cancellation and deadlines, edit/review retention, malformed-response rejection and persistent manual download.')


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--base', default='http://127.0.0.1:8095')
    run(parser.parse_args().base)
