"""Optional real-browser verification: requires Playwright and Chromium."""
import io
from pathlib import Path
import tempfile
from zipfile import ZipFile
from openpyxl import Workbook, load_workbook
from playwright.sync_api import sync_playwright

ROOT=Path(__file__).resolve().parents[1]

def run(base='http://127.0.0.1:8095'):
    with tempfile.TemporaryDirectory() as tmp, sync_playwright() as playwright:
        browser=playwright.chromium.launch(headless=True,executable_path='/usr/bin/chromium',args=['--no-sandbox'])
        page=browser.new_page(viewport={'width':1440,'height':1050},locale='vi-VN')
        errors=[];page.on('pageerror',lambda error:errors.append(str(error)))
        page.goto(base)
        page.locator('input[name=data]').set_input_files(str(ROOT/'samples/data-mau.xlsx'))
        page.locator('#analyze').click();page.locator('#workspace').wait_for(state='visible')
        assert page.locator('.sheet-card').count()==4
        assert page.locator('.selected:checked').count()==3
        assert '1.569.000' in page.locator('#total').inner_text()
        assert page.locator('[data-item="2"][data-item-field="note"]').input_value()=='Rau Nhà trường có'
        quantity=page.locator('[data-item="0"][data-item-field="quantity"]')
        quantity.fill('5');quantity.press('Tab')
        assert '1.653.000' in page.locator('#total').inner_text()
        page.locator('[data-row-toggle="1"]').click()
        assert '1.576.000' in page.locator('#total').inner_text()
        page.locator('[data-row-toggle="1"]').click()
        page.locator('[data-action=undo]').click()
        assert '1.569.000' in page.locator('#total').inner_text()
        page.locator('.bulk-panel summary').click()
        page.locator('#bulk-address').fill('Địa chỉ kiểm thử')
        page.locator('#bulk-date').fill('2026-10-02')
        page.locator('#apply-bulk').click()
        with page.expect_download() as info:page.locator('#export').click()
        output=Path(tmp)/'output.xlsx';info.value.save_as(output)
        workbook=load_workbook(output)
        assert len(workbook.worksheets)==3
        assert all(s['B9'].value=='Địa chỉ kiểm thử' for s in workbook)
        assert all(s['E6'].value.strftime('%Y-%m-%d')=='2026-10-02' for s in workbook)
        assert sum(s['H27'].value for s in workbook)==1569000
        page.locator('#export-format').select_option('zip')
        with page.expect_download() as info:page.locator('#export').click()
        output=Path(tmp)/'output.zip';info.value.save_as(output)
        with ZipFile(output) as archive:assert len(archive.namelist())==3
        page.locator('[data-tab=source]').click();assert page.locator('.preview-table').count()==1
        page.locator('[data-tab=mapping]').click()
        page.locator('[data-action=remap]').click()
        page.wait_for_function("document.getElementById('status').textContent.includes('Đã đọc lại')")
        assert page.locator('.selected:checked').count()==3
        page.locator('#sheet-search').fill('An Hà');assert page.locator('.sheet-card').count()==1
        page.locator('#sheet-search').fill('')
        page.locator('#data-file').set_input_files(str(ROOT/'samples/data-da-dang.xlsx'))
        page.locator('#analyze').click()
        page.wait_for_function("document.getElementById('source-filename').textContent.includes('data-da-dang')")
        assert page.locator('.selected:checked').count()==3
        page.locator('[data-open-sheet="1"]').click()
        assert 'dòng 4–5' in page.locator('#sheet-detail').inner_text()
        page.locator('[data-open-sheet="3"]').click()
        assert 'không có tiêu đề' in page.locator('#sheet-detail').inner_text().lower()
        # An invalid numeric draft remains visible when leaving/returning to the sheet.
        invalid=page.locator('[data-item="0"][data-item-field="quantity"]')
        invalid.fill('abc');page.locator('[data-tab=source]').click();page.locator('[data-tab=items]').click()
        assert page.locator('[data-item="0"][data-item-field="quantity"]').input_value()=='abc'
        page.locator('[data-action=undo]').click()
        # Remapping to a later customer must refresh automatically detected metadata.
        wb=Workbook();sheet=wb.active;sheet.title='Hai khách'
        rows=[['Khách hàng: Khách A'],['Ngày xuất: 01/10/2026'],['Địa chỉ: Địa chỉ A'],['Tên hàng','Số lượng','Đơn giá','Thành tiền'],['Gạo A',2,10000,20000],['Tổng cộng',None,None,20000],[],[],['Khách hàng: Khách B'],['Ngày xuất: 02/10/2026'],['Địa chỉ: Địa chỉ B'],['Tên hàng','Số lượng','Đơn giá','Thành tiền'],['Gạo B',3,20000,60000]]
        for row in rows:sheet.append(row)
        mixed_file=Path(tmp)/'data-hai-khach.xlsx';wb.save(mixed_file)
        page.locator('#data-file').set_input_files(str(mixed_file));page.locator('#analyze').click()
        page.wait_for_function("document.getElementById('source-filename').textContent.includes('data-hai-khach')")
        assert page.locator('[data-field=customer]').input_value()=='Khách A'
        page.locator('[data-tab=mapping]').click();page.locator('#mapping-header').fill('12')
        page.locator('[data-action=remap]').click()
        page.wait_for_function("document.getElementById('status').textContent.includes('Đã đọc lại')")
        assert page.locator('[data-field=customer]').input_value()=='Khách B'
        assert page.locator('[data-field=date]').input_value()=='2026-10-02'
        assert page.locator('[data-field=address]').input_value()=='Địa chỉ B'
        # A long sheet paginates and item search keeps the full export selection.
        wb=Workbook();sheet=wb.active;sheet.title='Bảng dài'
        sheet.append(['Khách hàng: Bếp ăn Minh An']);sheet.append(['Ngày 01 tháng 10 năm 2026'])
        sheet.append(['Tên hàng','Số lượng','Đơn giá','Thành tiền'])
        for index in range(1,84):sheet.append([f'Hàng {index}',1,2000,2000])
        long_file=Path(tmp)/'data-bang-dai.xlsx';wb.save(long_file)
        page.locator('#data-file').set_input_files(str(long_file));page.locator('#analyze').click()
        page.wait_for_function("document.getElementById('source-filename').textContent.includes('data-bang-dai')")
        assert page.locator('.item-table tbody tr').count()==50
        page.locator('[data-page="1"]').click();assert page.locator('.item-table tbody tr').count()==33
        page.locator('#item-search').fill('Hàng 83');assert page.locator('.item-table tbody tr').count()==1
        assert '166.000' in page.locator('#total').inner_text()
        # Original sample screenshot and mobile overflow check.
        page.locator('#data-file').set_input_files(str(ROOT/'samples/data-mau.xlsx'))
        page.locator('#analyze').click()
        page.wait_for_function("document.getElementById('source-filename').textContent.includes('data-mau')")
        for field in ['customer','recipient','date','address']:page.locator('#bulk-'+field).fill('')
        page.locator('.bulk-panel').evaluate('(element)=>element.open=false')
        page.locator('#export-format').select_option('xlsx')
        page.screenshot(path=str(ROOT/'samples/giao-dien-v3.png'),full_page=True)
        page.set_viewport_size({'width':390,'height':844})
        assert page.evaluate('document.documentElement.scrollWidth <= window.innerWidth')
        page.screenshot(path=str(ROOT/'samples/giao-dien-mobile-v3.png'),full_page=True)
        assert not errors,errors
        browser.close()
        print('Browser passed: upload, recognition, row edit/exclusion/undo, bulk metadata, workbook/ZIP, source/remap, search, split headers, inferred table, draft preservation, pagination, mobile.')

if __name__=='__main__':run()
