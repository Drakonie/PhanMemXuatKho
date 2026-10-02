import io
import sys
import unittest
from pathlib import Path
from decimal import Decimal
from datetime import date
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from engine import analyze, export, number

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = (ROOT/'templates/Template.xlsx').read_bytes()


def workbook(rows, title='Chi tiết'):
    wb=Workbook();wb.active.title=title
    for row in rows: wb.active.append(row)
    return wb


def data(wb):
    output=io.BytesIO();wb.save(output);return output.getvalue()


def simple(items=None):
    return workbook([
        ['Khách hàng: Công ty Minh Anh'],['Ngày xuất:',date(2026,10,1)],
        ['Tên hàng','ĐVT','Số lượng','Đơn giá','Thành tiền','Ghi chú'],
        *(items or [['Rau xanh','kg',2,10000,20000,'']]),
    ])


class RecognitionTests(unittest.TestCase):
    def test_english_aliases_shuffled_and_source_evidence(self):
        wb=workbook([
            ['Customer: Acme'],['Date: 2026-10-01'],['Address: 12 Market Street'],['Voucher no: XK-42'],
            ['Remarks','Unit price (VND)','Item description','QTY','Line total','UOM','SKU'],
            ['Fresh',15000,'Green beans',3,45000,'kg','GB01']
        ])
        result=analyze(data(wb))[0]
        self.assertEqual(result['mapping'],{'name':'C','quantity':'D','price':'B','amount':'E','unit':'F','code':'G','note':'A'})
        self.assertEqual((result['customer'],result['date'],result['address'],result['voucher']),('Acme','2026-10-01','12 Market Street','XK-42'))
        self.assertEqual(result['total'],45000)
        self.assertGreaterEqual(result['confidence'],90)
        self.assertEqual(result['preview_rows'][-1]['cells']['C'],'Green beans')
        self.assertTrue(result['detection_notes'])
        self.assertEqual(result['errors'],[])

    def test_split_two_row_labels_without_merges(self):
        wb=workbook([
            ['Khách hàng: Công ty Minh Anh'],['Ngày 01 tháng 10 năm 2026'],
            ['Tên','Số','Đơn','Thành'],['hàng','lượng','giá','tiền'],['Rau',2,10000,20000]
        ])
        result=analyze(data(wb))[0]
        self.assertEqual((result['header'],result['header_span']),(3,2))
        self.assertEqual(result['items'][0]['row'],5)
        self.assertEqual(result['date'],'2026-10-01')
        self.assertEqual(result['errors'],[])

    def test_merged_quantity_group_prefers_actual_issued(self):
        wb=workbook([
            ['Khách hàng: Công ty Minh Anh'],['Ngày 01 tháng 10 năm 2026'],
            ['Tên hàng','Số lượng',None,'Đơn giá','Thành tiền'],
            [None,'Yêu cầu','Thực xuất',None,None],['Rau',5,2,10000,20000]
        ])
        for area in ('A3:A4','B3:C3','D3:D4','E3:E4'): wb.active.merge_cells(area)
        result=analyze(data(wb))[0]
        self.assertEqual((result['header'],result['header_span']),(3,2))
        self.assertEqual(result['mapping']['quantity'],'C')
        self.assertEqual(result['items'][0]['quantity'],2)
        self.assertEqual(result['errors'],[])

    def test_multiple_tables_with_shuffled_second_header(self):
        wb=simple();s=wb.active
        s.append(['Tổng cộng',None,None,None,20000]);s.append([])
        s.append(['Đơn giá','Mặt hàng','Thành tiền','SL']);s.append([5000,'Cà rốt',15000,3]);s.append(['Tổng cộng',None,15000])
        result=analyze(data(wb))[0]
        self.assertEqual(len(result['tables']),2)
        self.assertEqual(result['tables'][1]['mapping']['name'],'B')
        self.assertEqual([item['row'] for item in result['items']],[4,8])
        self.assertEqual(result['total'],35000)
        self.assertEqual(result['errors'],[])
        manual=analyze(data(wb),{'Chi tiết':{'header':3,'mapping':{'name':'A','unit':'B','quantity':'C','price':'D','amount':'E','note':'F'}}})[0]
        self.assertEqual(manual['total'],35000)
        self.assertEqual(len(manual['tables']),2)

    def test_different_customer_between_tables_blocks_merge(self):
        wb=simple();s=wb.active
        s.append(['Tổng cộng',None,None,None,20000]);s.append(['Khách hàng: Công ty Khác'])
        s.append(['Tên hàng','Số lượng','Đơn giá','Thành tiền']);s.append(['Cà rốt',3,5000,15000])
        result=analyze(data(wb))[0]
        self.assertEqual(len(result['items']),2)
        self.assertTrue(any('khách hàng khác' in error for error in result['errors']))
        self.assertTrue(any(issue.get('field')=='customer' and not issue['recoverable'] for issue in result['issues']))
        with self.assertRaises(ValueError): export(TEMPLATE,[result])

    def test_missing_name_and_invalid_price_do_not_disappear(self):
        result=analyze(data(simple([[None,'kg',2,10000,20000,''],['Rau','kg',2,'mười nghìn',20000,'']])))[0]
        self.assertEqual(result['items'],[])
        self.assertTrue(any('thiếu tên hàng' in error for error in result['errors']))
        self.assertTrue(any('Không đọc được đơn giá' in error for error in result['errors']))
        self.assertTrue(all(not issue['recoverable'] for issue in result['issues'] if issue['level']=='error'))

    def test_amount_mismatch_retains_row_and_marks_recoverable(self):
        result=analyze(data(simple([['Rau','kg',2,10000,15000,'']])))[0]
        self.assertEqual(len(result['items']),1)
        issue=next(issue for issue in result['issues'] if issue['level']=='error')
        self.assertEqual((issue['row'],issue['field'],issue['recoverable']),(4,'amount',True))

    def test_formula_round_iferror_sum_ranges_and_cross_sheet(self):
        wb=simple([['Rau','kg','=IFERROR(1/0;2)',10000,'=ROUND(C4*D4;0)',''],['Cá','kg',3,30000,'=ROUNDUP(C5*D5,0)','']])
        lookup=wb.create_sheet("Giá và số lượng")
        lookup.append([2,3,4]);lookup.append([5,6,7])
        wb.active['C5']="=SUM('Giá và số lượng'!A1:B1,1,-3)"
        result=analyze(data(wb))[0]
        self.assertEqual(result['total'],110000)
        self.assertEqual(result['errors'],[])
        self.assertTrue(any('công thức' in note for note in result['detection_notes']))

    def test_unsupported_and_circular_formulas_block(self):
        wb=simple();wb.active['E4']='=VLOOKUP(A4,A1:F3,2,FALSE)'
        result=analyze(data(wb))[0]
        self.assertTrue(result['errors']);self.assertEqual(result['items'],[])
        wb=simple();wb.active['C4']='=C4+1'
        result=analyze(data(wb))[0]
        self.assertTrue(any('Công thức vòng' in error for error in result['errors']))

    def test_ambiguous_text_quantity_requires_explicit_review(self):
        result=analyze(data(simple([['Rau','kg','2.500',10000,25000000,'']])))[0]
        self.assertTrue(result['requires_review'])
        self.assertFalse(result['selected'])
        self.assertTrue(any('phân cách chưa rõ' in warning for warning in result['warnings']))
        self.assertLessEqual(result['confidence'],65)
        self.assertIsNone(number('1.00.000'))
        self.assertIsNone(number('1 2'))
        self.assertEqual(number('1 234,5 đ'),Decimal('1234.5'))
        self.assertEqual(number('(1.234,50 đ)'),Decimal('-1234.50'))

    def test_goods_names_and_notes_are_not_footer_or_headers(self):
        result=analyze(data(simple([
            ['Mặt hàng cá hồi','kg',2,10000,20000,'Total 20.000'],
            ['Tổng hợp rau sạch','kg',3,5000,15000,'Người nhận hàng xác nhận'],
            ['Rau cải','kg',1,5000,5000,'Người lập phiếu: Linh'],
        ])))[0]
        self.assertEqual(len(result['items']),3)
        self.assertEqual(result['total'],40000)
        self.assertEqual(result['errors'],[])

    def test_duplicate_quantity_headers_need_manual_selection(self):
        wb=workbook([['Tên hàng','SL','SL','Đơn giá','Thành tiền'],['Rau',2,3,10000,30000]])
        result=analyze(data(wb))[0]
        self.assertTrue(any('nhiều cột' in error for error in result['errors']))
        manual=analyze(data(wb),{'Chi tiết':{'header':1,'mapping':{'name':'A','quantity':'C','price':'D','amount':'E'}}})[0]
        self.assertEqual(manual['errors'],[])
        self.assertEqual(manual['total'],30000)

    def test_headerless_inference_is_a_proposal(self):
        wb=workbook([['Rau xanh',2,10000,20000],['Cá tươi',3,30000,90000]])
        result=analyze(data(wb))[0]
        self.assertEqual(result['detection_method'],'inferred')
        self.assertEqual(result['total'],110000)
        self.assertEqual(result['errors'],[])
        self.assertFalse(result['selected'])
        self.assertTrue(result['requires_review'])
        self.assertEqual(result['header'],0)

    def test_dotted_abbreviations_and_subtotal_do_not_drop_later_goods(self):
        wb=workbook([
            ['Tên hàng','Đ.V.T.','S.L.','Đ.G.','Thành tiền'],['Rau','kg',2,10000,20000],
            ['Subtotal',None,None,None,20000],['Cá','kg',3,30000,90000],['Tổng cộng',None,None,None,110000]
        ])
        result=analyze(data(wb))[0]
        self.assertEqual(result['mapping']['unit'],'B')
        self.assertEqual(result['total'],110000)
        self.assertEqual(len(result['items']),2)
        self.assertTrue(result['requires_review'])
        self.assertFalse(result['selected'])
        self.assertEqual(result['errors'],[])

    def test_missing_name_after_subtotal_blocks_instead_of_disappearing(self):
        wb=workbook([
            ['Tên hàng','Số lượng','Đơn giá','Thành tiền'],['Rau',2,10000,20000],
            ['Tổng cộng',None,None,20000],[None,3,5000,15000],['Tổng cộng',None,None,35000]
        ])
        result=analyze(data(wb))[0]
        self.assertEqual(result['total'],20000)
        self.assertTrue(any('Dòng 4' in error and 'thiếu tên hàng' in error for error in result['errors']))
        issue=next(issue for issue in result['issues'] if issue['level']=='error')
        self.assertEqual((issue['row'],issue['field'],issue['recoverable']),(4,'name',False))
        with self.assertRaises(ValueError): export(TEMPLATE,[{**result,'selected':True,'customer':'ABC','date':'2026-10-01'}])

    def test_changed_customer_after_subtotal_without_header_blocks(self):
        wb=simple();s=wb.active
        s.append(['Tổng cộng',None,None,None,20000]);s.append(['Khách hàng: Công ty Khác'])
        s.append(['Cá','kg',3,5000,15000])
        result=analyze(data(wb))[0]
        self.assertEqual(result['total'],35000)
        self.assertTrue(any(issue.get('field')=='customer' and issue['level']=='error' for issue in result['issues']))

    def test_later_table_customer_when_first_is_unknown_requires_split(self):
        wb=workbook([
            ['Tên hàng','Số lượng','Đơn giá','Thành tiền'],['Rau',2,10000,20000],['Tổng cộng',None,None,20000],
            ['Khách hàng: Công ty Minh Anh'],['Tên hàng','Số lượng','Đơn giá','Thành tiền'],['Cá',3,5000,15000]
        ])
        result=analyze(data(wb))[0]
        self.assertEqual(len(result['items']),2)
        self.assertTrue(any(issue.get('field')=='customer' and issue['level']=='error' for issue in result['issues']))

    def test_unrelated_numeric_sheet_is_not_selected(self):
        wb=workbook([['Nhân sự',2026,100,3],['Thời gian',2025,150,4]])
        result=analyze(data(wb))[0]
        self.assertFalse(result['selected']);self.assertEqual(result['items'],[])

    def test_hidden_and_summary_are_unselected(self):
        wb=simple();hidden=wb.copy_worksheet(wb.active);hidden.title='Sheet ẩn';hidden.sheet_state='hidden'
        summary=wb.copy_worksheet(wb.active);summary.title='Báo cáo tổng hợp'
        sheets=analyze(data(wb))
        self.assertTrue(sheets[0]['selected'])
        self.assertFalse(sheets[1]['selected']);self.assertEqual(sheets[1]['sheet_state'],'hidden')
        self.assertFalse(sheets[2]['selected']);self.assertTrue(sheets[2]['summary'])

    def test_formatted_trailing_rows_and_late_headers(self):
        wb=simple();wb.active['A1048576'].fill=PatternFill('solid',fgColor='FFFFFF')
        result=analyze(data(wb))[0]
        self.assertEqual(len(result['items']),1)
        wb=Workbook();s=wb.active
        s.cell(10001,1,'Tên hàng');s.cell(10001,2,'SL');s.cell(10001,3,'Đơn giá');s.cell(10001,4,'Thành tiền')
        for col,value in enumerate(['Rau',2,10000,20000],1):s.cell(10002,col,value)
        result=analyze(data(wb))[0]
        self.assertEqual(result['header'],10001)
        self.assertEqual(result['total'],20000)

    def test_long_names_and_notes_are_wrapped_in_export(self):
        name='Thịt bò nhập khẩu tươi đóng gói đủ định lượng giao nhà trường ' * 2
        note='Giao hàng trước giờ ăn trưa, bảo quản lạnh và kiểm tra chất lượng ' * 2
        result=analyze(data(simple([[name,'kg',2,10000,20000,note]])))[0]
        sheet=load_workbook(io.BytesIO(export(TEMPLATE,[result]))).active
        self.assertTrue(sheet['B12'].alignment.wrap_text)
        self.assertTrue(sheet['I12'].alignment.wrap_text)
        self.assertGreater(sheet.row_dimensions[12].height,23.1)
        self.assertEqual(str(sheet.page_setup.paperSize),str(sheet.PAPERSIZE_A4))


if __name__=='__main__':unittest.main()
