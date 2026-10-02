from __future__ import annotations
import ast
import calendar
import copy
import io
import math
import re
import unicodedata
from datetime import datetime, date
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP, ROUND_UP, ROUND_DOWN
from dependencies import configure_vendor

configure_vendor()

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.utils.cell import range_boundaries
from openpyxl.workbook.properties import CalcProperties


def normalize(value):
    text = unicodedata.normalize('NFKC', str(value if value is not None else ''))
    text = re.sub(r'[\u200b-\u200d\ufeff]', '', text).lower().replace('đ', 'd')
    return re.sub(r'\s+', ' ', ''.join(c for c in unicodedata.normalize('NFD', text) if not unicodedata.combining(c))).strip()


ALIASES = {
    'name': ['ten hang', 'ten hang hoa', 'ten san pham', 'hang', 'hang hoa', 'mat hang', 'ten vat tu', 'ten thuc pham', 'san pham',
             'ten nguyen lieu', 'nguyen lieu', 'noi dung hang hoa', 'ten hang hoa dich vu', 'item name', 'product name', 'item description', 'description', 'product', 'item', 'goods'],
    'quantity': ['so luong', 'sl', 'thuc xuat', 'sl xuat', 'so luong xuat', 'khoi luong', 'sl thuc xuat', 'so luong thuc xuat', 'quantity', 'qty', 'qty issued', 'issued quantity', 'weight', 'net weight', 'thuc te', 'sl thuc te', 'so luong thuc te', 'actual quantity'],
    'price': ['don gia', 'gia ban', 'gia', 'dg', 'don gia ban', 'don gia xuat', 'gia xuat', 'unit price', 'price', 'rate', 'unit cost'],
    'amount': ['thanh tien', 'thanh tien cong no', 'tong tien', 'tien hang', 'gia tri', 'amount', 'line total', 'total amount', 'total price', 'total', 'value'],
    'unit': ['dvt', 'don vi tinh', 'don vi', 'uom', 'unit', 'units', 'unit of measure'],
    'code': ['ma hang', 'ma hang hoa', 'ma san pham', 'ma vat tu', 'ma nguyen lieu', 'sku', 'item code', 'product code', 'code'],
    'note': ['ghi chu', 'dien giai', 'note', 'notes', 'remark', 'remarks'],
}


def header_scores(value):
    """Match whole labels and optional modifiers, never arbitrary substrings of item names."""
    text = normalize(value)
    text = re.sub(r'[.\-_]+', ' ', text)
    text = re.sub(r'\s+', ' ', text).strip(' :/\n')
    compact = re.sub(r'[\s./]', '', text)
    if compact in ('sl','dg','dvt','qty','uom','sku'): text = compact
    # Currency and unit suffixes are decorations, not competing header meanings.
    clean = re.sub(r'\([^)]*\)|\[[^]]*\]', ' ', text)
    clean = re.sub(r'\s+', ' ', clean).strip(' :/')
    clean_compact = clean.replace(' ', '')
    scores = {}
    for field, aliases in ALIASES.items():
        best = 0
        for alias in aliases:
            if clean == alias or text == alias or (len(alias.replace(' ', '')) >= 4 and clean_compact == alias.replace(' ', '')):
                best = max(best, 10 + len(alias.split()))
            elif len(alias) > 3 and re.match(r'^' + re.escape(alias) + r'(?:\s|/|:)', clean):
                suffix = clean[len(alias):].strip(' /:')
                # Metadata such as "Đơn vị mua" and "Tên hàng: thịt" is not a header.
                if suffix and len(suffix.split()) <= 5 and not re.search(r'khach|nguoi|mua|customer|buyer', suffix):
                    best = max(best, 7 + len(alias.split()))
        if field == 'quantity' and best and re.search(r'thuc xuat|thuc te|actual|issued', text):
            best += 4
        if best:
            scores[field] = best
    return scores



# Split labels may occupy up to three rows. Only known header fragments may
# extend a header; an incomplete goods row must remain a visible data error.
HEADER_FRAGMENTS = {" ".join(words[start:end])
                    for aliases in ALIASES.values() for alias in aliases
                    for words in [alias.split()] for start in range(len(words))
                    for end in range(start+1,len(words)+1)}
HEADER_FRAGMENTS.update({'yeu cau', 'theo chung tu', 'thuc', 'xuat', 'vnd', 'vnd/kg',
                         'vnd kg', 'dong', 'd', 'kg', 'g', 'lit', 'l',
                         'usd', 'so tt', 'stt', 'quy cach', 'pham chat', 'hoa'})


def header_fragment(value):
    if not isinstance(value,str) or value.startswith('='):
        return False
    text = re.sub(r'[().,\[\]:_\-]+', ' ', normalize(value))
    text = re.sub(r'\s+', ' ', text).strip()
    return text in HEADER_FRAGMENTS or any(text.replace(' ', '') == alias.replace(' ', '')
                                         for aliases in ALIASES.values() for alias in aliases)


def field_of(value):
    scores = header_scores(value)
    return max(scores, key=scores.get) if scores else None


def number(value):
    """Read finite numeric cells and common Vietnamese/international number strings."""
    if isinstance(value, bool) or value is None or str(value).strip() == '':
        return None
    if isinstance(value, (int, float, Decimal)):
        try:
            result = Decimal(str(value))
        except InvalidOperation:
            return None
    else:
        text = str(value).strip().replace('\u2212', '-')
        negative = text.startswith('(') and text.endswith(')')
        if negative:
            text = text[1:-1].strip()
        text = re.sub(r'(?i)^(?:vnd|vnđ|đồng|₫|đ|dong)\s*|\s*(?:vnd|vnđ|đồng|₫|đ|dong)$', '', text).strip()
        if negative and text.startswith(('+', '-')):
            return None
        if re.search(r'\s', text):
            if not re.fullmatch(r'[+-]?\d{1,3}(?:[ \u00a0\u202f]\d{3})+(?:[.,]\d+)?', text):
                return None
            text = re.sub(r'\s', '', text)
        if not re.fullmatch(r'[+-]?\d+(?:[.,]\d+)*', text):
            return None
        if ',' in text and '.' in text:
            decimal_sep = ',' if text.rfind(',') > text.rfind('.') else '.'
            grouping_sep = '.' if decimal_sep == ',' else ','
            integer, fractional = text.rsplit(decimal_sep, 1)
            if decimal_sep in integer or not re.fullmatch(r'[+-]?\d{1,3}(?:' + re.escape(grouping_sep) + r'\d{3})+', integer):
                return None
            text = integer.replace(grouping_sep, '') + '.' + fractional
        elif ',' in text or '.' in text:
            sep = ',' if ',' in text else '.'
            parts = text.split(sep)
            if len(parts) > 2:
                if not re.fullmatch(r'[+-]?\d{1,3}(?:' + re.escape(sep) + r'\d{3})+', text):
                    return None
                text = ''.join(parts)
            elif len(parts[1]) == 3 and len(parts[0].lstrip('+-')) <= 3 and parts[0].lstrip('+-') != '0':
                text = ''.join(parts)
            elif sep == ',':
                text = text.replace(',', '.')
        try:
            result = Decimal(text)
        except InvalidOperation:
            return None
        if negative:
            result = -result
    return result if result.is_finite() else None


class Values:
    """Use saved Excel results, otherwise evaluate a small explicit formula language."""
    def __init__(self, raw, cached):
        self.raw, self.cached, self.memo, self.active = raw, cached, {}, set()
        self.evaluated = set()

    def get(self, sheet, coord):
        coord = coord.replace('$', '').upper()
        key = (sheet, coord)
        if key in self.memo:
            return self.memo[key]
        cell = self.raw[sheet][coord]
        value = cell.value
        if cell.data_type == 'e':
            raise ValueError(f'Ô {coord} có lỗi Excel {value}')
        if cell.data_type != 'f' or not isinstance(value, str) or not value.startswith('='):
            return value
        saved = self.cached[sheet][coord]
        if saved.data_type == 'e':
            raise ValueError(f'Ô {coord} có lỗi Excel {saved.value}')
        if saved.value is not None:
            return saved.value
        if key in self.active or len(self.active) > 80:
            raise ValueError('Công thức vòng hoặc quá phức tạp')
        if len(value) > 20000:
            raise ValueError('Công thức quá dài')
        self.active.add(key)
        try:
            expression = value[1:].replace(';', ',').replace('^', '**')
            # Keep quoted literal strings separate so A1 inside text is not a reference.
            pieces = re.split(r'("(?:[^"]|"")*")', expression)
            ref_pattern = re.compile(r"(?<![\w.])(?:'((?:[^']|'')+)'!|([\w.]+)!)?(\$?[A-Za-z]{1,3}\$?\d+)(?::(\$?[A-Za-z]{1,3}\$?\d+))?(?![\w.(])")
            def replace_reference(match):
                target = (match.group(1) or match.group(2) or sheet).replace("''", "'")
                first = match.group(3).replace('$', '').upper()
                last = match.group(4)
                if last:
                    return f'_RANGE({target!r},{(first + ":" + last.replace("$", "").upper())!r})'
                return f'_REF({target!r},{first!r})'
            for index in range(0, len(pieces), 2):
                pieces[index] = ref_pattern.sub(replace_reference, pieces[index])
                pieces[index] = pieces[index].replace('<>', '!=')
                pieces[index] = re.sub(r'(?<![<>=!])=(?!=)', '==', pieces[index])
            expression = ''.join(pieces)
            tree = ast.parse(expression, mode='eval')
            if sum(1 for _ in ast.walk(tree)) > 5000:
                raise ValueError('Công thức quá phức tạp')
            for expression_node in ast.walk(tree):
                if isinstance(expression_node, ast.BinOp) and isinstance(expression_node.op, ast.Pow):
                    if any(isinstance(child, ast.Pow) for branch in (expression_node.left, expression_node.right) for child in ast.walk(branch)):
                        raise ValueError('Công thức có số mũ lồng nhau; hãy tính lại và lưu kết quả trong Excel')
                if isinstance(expression_node, ast.UnaryOp) and isinstance(expression_node.op, ast.USub):
                    if any(isinstance(child, ast.Pow) for child in ast.walk(expression_node.operand)):
                        raise ValueError('Công thức có dấu âm kèm số mũ; hãy tính lại và lưu kết quả trong Excel')

            def scalar(value):
                if isinstance(value, bool):
                    return Decimal(int(value))
                parsed = number(value)
                if parsed is None:
                    raise ValueError('Công thức tham chiếu ô không phải số')
                return parsed

            def flatten(args):
                for arg in args:
                    yield from arg if isinstance(arg, list) else [arg]

            def evaluate(node):
                if isinstance(node, ast.Expression):
                    return evaluate(node.body)
                if isinstance(node, ast.Constant):
                    if type(node.value) in (int, float):
                        return Decimal(str(node.value))
                    if type(node.value) in (str, bool):
                        return node.value
                if isinstance(node, ast.Name) and node.id.upper() in ('TRUE', 'FALSE'):
                    return node.id.upper() == 'TRUE'
                if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.UAdd, ast.USub)):
                    return scalar(evaluate(node.operand)) * (-1 if isinstance(node.op, ast.USub) else 1)
                if isinstance(node, ast.BinOp) and isinstance(node.op, (ast.Add, ast.Sub, ast.Mult, ast.Div, ast.Pow)):
                    left, right = scalar(evaluate(node.left)), scalar(evaluate(node.right))
                    if isinstance(node.op, ast.Add): return left + right
                    if isinstance(node.op, ast.Sub): return left - right
                    if isinstance(node.op, ast.Mult): return left * right
                    if isinstance(node.op, ast.Div): return left / right
                    if right != int(right) or abs(right) > 20: raise ValueError('Số mũ chưa hỗ trợ')
                    return left ** int(right)
                if isinstance(node, ast.Compare) and len(node.ops) == 1:
                    left, right = evaluate(node.left), evaluate(node.comparators[0])
                    operator = node.ops[0]
                    if isinstance(operator, ast.Eq): return left == right
                    if isinstance(operator, ast.NotEq): return left != right
                    if isinstance(operator, ast.Lt): return left < right
                    if isinstance(operator, ast.LtE): return left <= right
                    if isinstance(operator, ast.Gt): return left > right
                    if isinstance(operator, ast.GtE): return left >= right
                if isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and not node.keywords:
                    function = node.func.id.upper()
                    if function == 'IFERROR' and len(node.args) == 2:
                        try: return evaluate(node.args[0])
                        except (ValueError, ArithmeticError, KeyError): return evaluate(node.args[1])
                    if function == 'IF' and len(node.args) == 3:
                        return evaluate(node.args[1] if evaluate(node.args[0]) else node.args[2])
                    args = [evaluate(arg) for arg in node.args]
                    if function == '_REF' and len(args) == 2:
                        result = self.get(args[0], args[1])
                        return Decimal(0) if result is None else result
                    if function == '_RANGE' and len(args) == 2:
                        a, b, c, d = range_boundaries(args[1])
                        if (c-a+1)*(d-b+1) > 100000:
                            raise ValueError('Vùng công thức quá lớn')
                        return [self.get(args[0], f'{get_column_letter(col)}{row}') for row in range(b,d+1) for col in range(a,c+1)]
                    if function in ('SUM', 'MIN', 'MAX', 'PRODUCT', 'AVERAGE'):
                        numeric = []
                        for argument, argument_value in zip(node.args, args):
                            reference = isinstance(argument, ast.Call) and isinstance(argument.func, ast.Name) and argument.func.id.upper() in ('_REF', '_RANGE')
                            if reference and argument.func.id.upper() == '_REF':
                                argument_value = self.get(evaluate(argument.args[0]), evaluate(argument.args[1]))
                            if reference or isinstance(argument_value, list):
                                for item in flatten([argument_value]):
                                    if type(item) in (int, float, Decimal):
                                        parsed = number(item)
                                        if parsed is not None: numeric.append(parsed)
                            else:
                                numeric.append(scalar(argument_value))
                        if function == 'SUM': return sum(numeric, Decimal(0))
                        if function == 'MIN': return min(numeric, default=Decimal(0))
                        if function == 'MAX': return max(numeric, default=Decimal(0))
                        if function == 'AVERAGE': return sum(numeric, Decimal(0)) / len(numeric)
                        result = Decimal(1)
                        for value in numeric: result *= value
                        return result if numeric else Decimal(0)
                    if function in ('ROUND', 'ROUNDUP', 'ROUNDDOWN') and len(args) == 2:
                        digits = scalar(args[1])
                        if digits != int(digits) or abs(digits) > 12:
                            raise ValueError('Số chữ số ROUND không hợp lệ')
                        rounding = {'ROUND': ROUND_HALF_UP, 'ROUNDUP': ROUND_UP, 'ROUNDDOWN': ROUND_DOWN}[function]
                        return scalar(args[0]).quantize(Decimal(1).scaleb(-int(digits)), rounding=rounding)
                    if function == 'ABS' and len(args) == 1: return abs(scalar(args[0]))
                    if function == 'INT' and len(args) == 1: return scalar(args[0]).to_integral_value(rounding='ROUND_FLOOR')
                raise ValueError('Công thức chưa hỗ trợ; hãy tính lại và lưu file trong Excel')

            result = evaluate(tree)
            if isinstance(result, Decimal) and not result.is_finite():
                raise ValueError('Công thức trả về số không hợp lệ')
            self.memo[key] = result
            self.evaluated.add(key)
            return result
        except (SyntaxError, TypeError, KeyError, ArithmeticError) as exc:
            raise ValueError(f'Không tính được công thức tại {coord}: {exc}') from exc
        finally:
            self.active.remove(key)


CUSTOMER_LABEL = r'(?:ten khach hang|khach hang|ho ten nguoi mua hang|ten don vi mua|don vi mua hang|don vi mua|nguoi mua hang|buyer|customer(?: name)?|ten don vi)'
DATE_LABEL = r'(?:ngay xuat(?: kho)?|ngay giao(?: hang)?|ngay lap|ngay|date|delivery date|issue date)'


def metadata(sheet, before, after=1):
    result = {'customer': '', 'date': '', 'address': '', 'voucher': '', 'recipient': '', 'content': ''}
    candidates = []
    def adjacent(cell):
        for column in range(cell.column + 1, min(sheet.max_column, cell.column + 12) + 1):
            value = sheet.cell(cell.row, column).value
            if value is not None and not (isinstance(value, str) and value.startswith('=')):
                return value
        return ''
    labels = {
        'customer': CUSTOMER_LABEL,
        'address': r'(?:dia chi(?: giao hang)?|delivery address|address)',
        'voucher': r'(?:so phieu(?: xuat(?: kho)?)?|so chung tu|voucher(?: no)?|invoice(?: no)?|document no)',
        'recipient': r'(?:ho ten nguoi nhan(?: hang)?|nguoi nhan(?: hang)?|recipient|receiver)',
        'content': r'(?:ly do xuat(?: kho)?|noi dung|reason|purpose)',
    }
    for row in sheet.iter_rows(min_row=max(1, after), max_row=max(1, before-1), max_col=min(sheet.max_column,100)):
        for cell in row:
            value = cell.value
            if isinstance(value, (datetime, date)):
                result['date'] = value.strftime('%Y-%m-%d')
                continue
            if not isinstance(value, str) or value.startswith('='):
                continue
            norm = normalize(value)
            for key, label in labels.items():
                match = re.match(r'^(' + label + r')\s*(?:[:：]\s*)?(.*)$', norm)
                if match:
                    # Preserve original Vietnamese text rather than returning normalized labels.
                    original = re.split(r'[:：]', value, maxsplit=1)
                    if len(original) == 2 and original[1].strip():
                        result[key] = original[1].strip()
                    elif match[2].strip():
                        words = len(match[1].split())
                        result[key] = ' '.join(value.strip().split()[words:]).strip(' :：')
                    else:
                        result[key] = str(adjacent(cell)).strip()
            day_match = re.search(r'ngay\s+(\d{1,2})\s+thang\s+(\d{1,2})\s+nam\s+(\d{4})', norm)
            if not day_match:
                day_match = re.search(r'\b(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{4})\b', norm)
            if day_match:
                try: result['date'] = date(int(day_match[3]), int(day_match[2]), int(day_match[1])).isoformat()
                except ValueError: pass
            else:
                iso = re.search(r'\b(\d{4})-(\d{1,2})-(\d{1,2})\b', norm)
                if iso:
                    try: result['date'] = date(int(iso[1]), int(iso[2]), int(iso[3])).isoformat()
                    except ValueError: pass
                elif re.fullmatch(DATE_LABEL + r'\s*[:：]?', norm):
                    near = adjacent(cell)
                    if isinstance(near, (date, datetime)): result['date'] = near.strftime('%Y-%m-%d')
                    elif isinstance(near, str):
                        match = re.search(r'\b(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{4})\b', near)
                        if match:
                            try: result['date'] = date(int(match[3]), int(match[2]), int(match[1])).isoformat()
                            except ValueError: pass
            if re.search(r'^(?:truong\b|cong ty\b|nha hang\b|mam non\b|khach san\b)', norm):
                candidates.append(value.strip())
    if not result['customer'] and len(candidates) == 1:
        result['customer'] = candidates[0]
    return result


def customer_and_date(sheet, header):
    info = metadata(sheet, header)
    return info['customer'], info['date']


FOOTER = re.compile(r'^(?:tong cong(?: tien thanh toan)?|tong tien(?: thanh toan| hang)?|tong thanh tien|cong tien(?: hang)?|cong thanh tien|tong hop|subtotal|grand total|invoice total|amount due|total payable|total(?: amount)?|cong|tong)(?:\s*[:：].*|\s*\([^)]*\)|\s*)$')
SIGNATURE = re.compile(r'^(?:nguoi lap|nguoi nhan|nguoi giao|thu kho|ke toan|giam doc|ky ten|bang chu|so tien bang chu|prepared by|received by|approved by|signature|amount in words)(?:\b|:)')


def content_extent(sheet):
    populated = [cell for cell in sheet._cells.values() if cell.value is not None]
    return (max((cell.row for cell in populated),default=1), max((cell.column for cell in populated),default=1))


def header_candidates(sheet):
    """Collect one to three-row headers; merged cells carry their parent label down."""
    content_row,content_col = content_extent(sheet)
    limit_col = min(content_col, 150)
    limit_row = content_row
    possible_rows = set()
    for cell in list(sheet._cells.values()):
        if cell.column > limit_col or not isinstance(cell.value,str) or cell.value.startswith('='): continue
        text = normalize(cell.value)
        if header_scores(cell.value) or text in ('ten','so','don','thanh','khoi','unit'):
            possible_rows.add(cell.row)
            possible_rows.update(range(max(1,cell.row-2),cell.row))
    merged = {}
    for area in sheet.merged_cells.ranges:
        if area.min_row <= limit_row and area.min_col <= limit_col:
            value = sheet.cell(area.min_row, area.min_col).value
            if header_fragment(value):
                for row in range(area.min_row, min(area.max_row, limit_row) + 1):
                    for col in range(area.min_col, min(area.max_col, limit_col) + 1):
                        merged[(row, col)] = value
    candidates = []
    for row_index in sorted(possible_rows):
        row_values = [sheet.cell(row_index, col).value or merged.get((row_index, col)) for col in range(1,limit_col+1)]
        if sum(bool(header_scores(value)) for value in row_values) < 2:
            # A top row can contain only one horizontally merged parent header.
            split_possible = any(row_index+size-1<=limit_row and sum(bool(header_scores(' '.join(str(sheet.cell(row_index+offset,col).value or merged.get((row_index+offset,col)) or '') for offset in range(size)))) for col in range(1,limit_col+1)) >= 3 for size in (2,3))
            if not split_possible and not any((row_index, col) in merged for col in range(1,limit_col+1)):
                continue
        for span in (1, 2, 3):
            if row_index + span - 1 > limit_row: continue
            if span > 1:
                if not any(header_scores(value) or header_fragment(value) for value in row_values): continue
                extra_rows = [[sheet.cell(row_index+offset,col).value for col in range(1,limit_col+1) if sheet.cell(row_index+offset,col).value is not None] for offset in range(1,span)]
                if any(not row or not all(header_fragment(value) for value in row) for row in extra_rows): continue
            choices = {}
            second_fields = 0
            for col in range(1,limit_col+1):
                values = [sheet.cell(row_index+offset,col).value or merged.get((row_index+offset,col)) for offset in range(span)]
                # A data row with quantity values cannot be part of a header.
                if span > 1 and isinstance(values[-1], (int,float,Decimal,datetime,date)):
                    continue
                scores = {}
                for offset, value in enumerate(values):
                    for field, score in header_scores(value).items():
                        scores[field] = max(scores.get(field,0), score + (1 if span > 1 and offset > 0 else 0))
                        if offset > 0: second_fields += 1
                if span > 1 and all(value is None or isinstance(value,str) for value in values):
                    joined = ' '.join(str(v) for v in values if v)
                    for field, score in header_scores(joined).items():
                        scores[field] = max(scores.get(field,0),score+2)
                for field, score in scores.items():
                    choices.setdefault(field, []).append((score,col))
            mapping, ambiguous = {}, []
            occupied = set()
            for field in ('name','quantity','price','amount','unit','code','note'):
                ranked = sorted(choices.get(field,[]), reverse=True)
                ranked = [pair for pair in ranked if pair[1] not in occupied]
                if not ranked: continue
                best = ranked[0]
                if len(ranked)>1 and ranked[1][0] == best[0]:
                    ambiguous.append(field)
                mapping[field] = best[1]
                occupied.add(best[1])
            if span > 1 and not second_fields:
                # Joined split labels such as "Số" + "lượng" must still be considered.
                split_labels = any(header_scores(' '.join(str(sheet.cell(row_index+offset,col).value or '') for offset in range(span))) for col in range(1,limit_col+1))
                if not split_labels: continue
            if 'name' in mapping and 'quantity' in mapping and ('price' in mapping or 'amount' in mapping):
                # Only extend a complete header if row two adds meaning or is a genuine subheader.
                if span > 1:
                    nonempty = [sheet.cell(row_index+1,col).value for col in range(1,limit_col+1) if sheet.cell(row_index+1,col).value is not None]
                    if not nonempty or any(number(v) is not None or (isinstance(v,str) and v.startswith('=')) for v in nonempty):
                        continue
                    if second_fields == 0 and not split_labels: continue
                score = 76 + 4*int('price' in mapping) + 4*int('amount' in mapping) + min(9,3*len(set(mapping)-{'name','quantity','price','amount'}))
                candidates.append({'header':row_index,'header_span':span,'mapping':mapping,'confidence':min(96,score),'ambiguous':ambiguous,'method':'headers'})
    # Competing spans at the same header: choose the one with more semantic fields.
    unique = []
    for candidate in sorted(candidates,key=lambda c:(c['header'],-len(c['mapping']),-c['header_span'])):
        if unique and candidate['header'] <= unique[-1]['header']+unique[-1]['header_span']-1:
            continue
        unique.append(candidate)
    return unique


def infer_table(sheet, values):
    """Propose a table only when several rows support qty × price = amount."""
    limit_col = min(sheet.max_column,60)
    limit_row = min(sheet.max_row,500)
    row_stats = []
    for row in range(1,limit_row+1):
        texts,numbers = {},{}
        for col in range(1,limit_col+1):
            try: value = values.get(sheet.title,f'{get_column_letter(col)}{row}')
            except Exception: continue
            parsed = number(value)
            if parsed is not None and parsed >= 0: numbers[col] = parsed
            elif isinstance(value,str) and value.strip() and not value.startswith('=') and not field_of(value) and not FOOTER.match(normalize(value)) and not SIGNATURE.match(normalize(value)):
                if len(value.strip()) <= 200: texts[col] = value.strip()
        if texts and 3 <= len(numbers) <= 12:
            row_stats.append((row,texts,numbers))
    proposals = {}
    for row,texts,numbers in row_stats:
        for name_col in texts:
            for amount_col,amount in numbers.items():
                if amount <= 0: continue
                for qty_col,qty in numbers.items():
                    if qty_col == amount_col or qty <= 0: continue
                    for price_col,price in numbers.items():
                        if price_col in (amount_col,qty_col) or price <= 0: continue
                        if abs((qty*price).quantize(Decimal(1),rounding=ROUND_HALF_UP)-amount) <= 1:
                            key = (name_col,qty_col,price_col,amount_col)
                            proposals.setdefault(key,[]).append(row)
    if not proposals: return None
    ranked = sorted(proposals.items(),key=lambda pair:(len(pair[1]),pair[0][2]>pair[0][1]),reverse=True)
    key,rows = ranked[0]
    # Multiplication is symmetric, so price/quantity need additional evidence.
    name_col,qty_col,price_col,amount_col = key
    if len(rows) < 2: return None
    qty_values = [numbers[qty_col] for row,_,numbers in row_stats if row in rows]
    price_values = [numbers[price_col] for row,_,numbers in row_stats if row in rows]
    if not (max(qty_values) < min(price_values) and min(price_values) >= 100): return None
    # Multiple text columns (e.g. item code and item name) favor text with natural words.
    text_scores = {}
    for row,texts,_ in row_stats:
        if row in rows:
            for col,text in texts.items():
                text_scores[col] = text_scores.get(col,0) + int(' ' in text) + int(not re.fullmatch(r'[A-Za-z]{0,5}[-_]?\d+',text))
    if text_scores:
        best_name = max(text_scores,key=text_scores.get)
        if best_name != name_col:
            alternate = (best_name,qty_col,price_col,amount_col)
            if alternate in proposals and len(proposals[alternate]) == len(rows): name_col = best_name
    start = min(rows)
    return {'header':max(0,start-1),'header_span':1 if start>1 else 0,'mapping':{'name':name_col,'quantity':qty_col,'price':price_col,'amount':amount_col},'confidence':58,'ambiguous':[],'method':'inferred','start':start}


def safe_preview(sheet, header, span):
    content_row,content_col = content_extent(sheet)
    start = max(1,header-2)
    end = min(content_row,max(header+span+6,8))
    def display(value):
        if value is None: return ''
        if isinstance(value,(date,datetime)): return value.isoformat()
        return str(value)[:400]
    return [{'row':row,'cells':{get_column_letter(col):display(sheet.cell(row,col).value) for col in range(1,min(content_col,40)+1)}} for row in range(start,end+1)]


def analyze(source, overrides=None):
    raw = load_workbook(io.BytesIO(source),data_only=False)
    cached = load_workbook(io.BytesIO(source),data_only=True)
    values = Values(raw,cached)
    results = []
    for sheet in raw:
        content_row,content_col = content_extent(sheet)
        warnings, errors, issues, notes = [], [], [], []
        def issue(level,message,row=None,field=None,recoverable=False):
            (errors if level=='error' else warnings).append(message)
            record = {'level':level,'message':message,'recoverable':recoverable}
            if row is not None: record['row'] = row
            if field: record['field'] = field
            issues.append(record)
        override = (overrides or {}).get(sheet.title)
        if override:
            header = int(override['header'])
            if not 1 <= header <= sheet.max_row: raise ValueError('Dòng tiêu đề ngoài phạm vi sheet')
            span = int(override.get('header_span',1))
            if span not in (1,2,3) or header+span-1 > sheet.max_row: raise ValueError('Số dòng tiêu đề không hợp lệ')
            mapping = {key:column_index_from_string(str(value).strip().upper()) for key,value in override['mapping'].items() if value and key in ALIASES}
            if any(col<1 or col>min(sheet.max_column,150) for col in mapping.values()): raise ValueError('Cột đã chọn ngoài phạm vi sheet')
            if len(set(mapping.values())) != len(mapping): raise ValueError('Mỗi trường phải dùng một cột khác nhau')
            tables = [{'header':header,'header_span':span,'mapping':mapping,'confidence':95,'ambiguous':[],'method':'manual'}]
            tables.extend(candidate for candidate in header_candidates(sheet) if candidate['header'] >= header+span)
            notes.append('Cột và dòng tiêu đề do người dùng chỉ định.')
        else:
            tables = header_candidates(sheet)
            if not tables:
                inferred = infer_table(sheet,values)
                tables = [inferred] if inferred else []
        primary = tables[0] if tables else {'header':0,'header_span':1,'mapping':{},'confidence':0,'method':'none'}
        header, span, mapping = primary['header'],primary['header_span'],primary['mapping']
        summary = any(term in normalize(sheet.title) for term in ['tong hop','tong cong','summary','total','bao cao'])
        info = metadata(sheet,header or 1)
        result = {'sheet':sheet.title,'selected':bool(tables) and primary['method']!='inferred' and not summary and sheet.sheet_state=='visible',
                  'summary':summary,'sheet_state':sheet.sheet_state,'header':header,'header_span':span,
                  'mapping':{key:get_column_letter(col) for key,col in mapping.items()},
                  'columns':[get_column_letter(i) for i in range(1,min(content_col,150)+1)],
                  **info,'recipient':info['recipient'] or sheet.title,'items':[],'warnings':warnings,'errors':errors,'issues':issues,'total':0,
                  'confidence':primary['confidence'],'detection_method':primary['method'],'detection_notes':notes,
                  'preview_rows':safe_preview(sheet,header,span),'tables':[],'requires_review':False}
        results.append(result)
        if not tables or not all(key in mapping for key in ['name','quantity']) or not ('price' in mapping or 'amount' in mapping):
            result['selected'] = False
            issue('error','Chưa nhận diện đủ tên hàng, số lượng và đơn giá hoặc thành tiền. Hãy chọn dòng tiêu đề và cột thủ công.')
            continue
        if primary['method']=='inferred':
            notes.append('Đề xuất cột dựa trên nhiều dòng có số lượng × đơn giá khớp thành tiền; chưa có tiêu đề rõ ràng.')
            issue('warning','Bảng không có tiêu đề rõ ràng. Hãy kiểm tra cột và xác nhận chọn sheet trước khi xuất.')
            result['requires_review'] = True
        else:
            notes.append(f'Nhận diện {len(mapping)} trường tại dòng {header}' + (f'–{header+span-1}.' if span>1 else '.'))
        notes.append('Điểm nhận diện dựa trên tiêu đề và kiểm tra dòng hàng; không phải xác suất thống kê.')
        if len(tables)>1: notes.append(f'Tìm được {len(tables)} bảng/tiêu đề lặp trong sheet; gộp các dòng hàng hợp lệ.')
        if not info['customer']: issue('warning','Chưa tìm được tên khách hàng. Cần nhập trước khi xuất.')
        if not info['date']: issue('warning','Chưa tìm được ngày xuất. Cần nhập trước khi xuất.')
        if summary: issue('warning','Sheet tổng hợp được bỏ chọn mặc định để tránh xuất trùng với các sheet chi tiết.')
        if sheet.sheet_state!='visible': issue('warning','Sheet ẩn được bỏ chọn mặc định.')
        if not override and header > 1:
            for earlier_row in range(1,header):
                def prior_value(column):
                    try:
                        return values.get(sheet.title, f'{get_column_letter(column)}{earlier_row}')
                    except (ValueError, KeyError, ArithmeticError):
                        return sheet.cell(earlier_row,column).value
                earlier_name = prior_value(mapping['name'])
                earlier_numbers = {field: prior_value(mapping[field])
                                   for field in ('quantity','price','amount') if field in mapping}
                if sum(number(value) is not None for value in earlier_numbers.values()) >= 2 and not FOOTER.match(normalize(earlier_name)):
                    issue('error', f'Dòng {earlier_row}: có số liệu trước tiêu đề nhận diện. Hãy chọn lại tiêu đề/cột để đọc đủ các dòng hàng.', row=earlier_row)
                    result['requires_review'] = True
                    result['selected'] = False
        previous_end = 1
        for table_index,table in enumerate(tables):
            table_header,table_span,table_map = table['header'],table['header_span'],table['mapping']
            start = table.get('start',table_header+table_span)
            end = tables[table_index+1]['header']-1 if table_index+1<len(tables) else content_row
            table_record = {'header':table_header,'header_span':table_span,'mapping':{key:get_column_letter(col) for key,col in table_map.items()},'confidence':table['confidence'],'start':start,'end':end}
            result['tables'].append(table_record)
            if table_index:
                extra_info = metadata(sheet,table_header,previous_end)
                labels = {'customer':'khách hàng','date':'ngày xuất','voucher':'số phiếu','recipient':'người nhận','address':'địa chỉ'}
                for field in labels:
                    if extra_info[field] and (not info[field] or normalize(extra_info[field])!=normalize(info[field])):
                        label=labels[field]
                        if not info[field] and field not in ('customer','date','voucher'):
                            issue('warning', f'Bảng tại dòng {table_header} có {label} nhưng bảng đầu chưa rõ. Hãy kiểm tra và điền thông tin cho phiếu chung.', field=field)
                            result['requires_review'] = True
                            result['selected'] = False
                            continue
                        message=f'Bảng tại dòng {table_header} có {label} khác bảng đầu. Hãy tách các phiếu ra sheet riêng.' if info[field] else f'Bảng tại dòng {table_header} có {label} nhưng bảng đầu chưa rõ thông tin. Chưa đủ cơ sở gộp các bảng; hãy tách các phiếu ra sheet riêng.'
                        issue('error',message,field=field)
            last_item_row = start-1
            for ambiguous_field in table['ambiguous']:
                if ambiguous_field in ('name','quantity','price','amount'):
                    issue('error',f'Dòng {table_header}: có nhiều cột cùng phù hợp với trường {ambiguous_field}. Hãy chọn cột thủ công.',row=table_header,field=ambiguous_field)
            # Infer a note column only from short text accompanying real item rows.
            if not override and 'note' not in table_map:
                for col in range(max(table_map.values(),default=0)+1,min(sheet.max_column,150)+1):
                    if sheet.cell(table_header or start,col).value is not None: continue
                    found = []
                    for row in range(start,min(end,start+49)+1):
                        item_name = sheet.cell(row,table_map['name']).value
                        if FOOTER.match(normalize(item_name)): break
                        value = sheet.cell(row,col).value
                        if item_name and isinstance(value,str) and not value.startswith('=') and number(value) is None and len(value)<300:
                            found.append(row)
                    if found:
                        table_map['note'] = col
                        table_record['mapping']['note'] = get_column_letter(col)
                        if table_index==0: result['mapping']['note'] = get_column_letter(col)
                        issue('warning',f'Nhận diện ghi chú ở cột {get_column_letter(col)} dù cột chưa có tiêu đề; hãy kiểm tra.')
                        break
            footer_reached = False
            for row_index in range(start,end+1):
                def read(field):
                    if field not in table_map: return None
                    return values.get(sheet.title,f'{get_column_letter(table_map[field])}{row_index}')
                try:
                    name = read('name')
                    label_values = [sheet.cell(row_index,col).value for col in range(1,table_map['name']+1)]
                    normalized = [normalize(value) for value in label_values if isinstance(value,str) and not value.startswith('=')]
                    named_goods = isinstance(name,str) and bool(name.strip()) and not FOOTER.match(normalize(name))
                    if any(FOOTER.match(text) for text in normalized) and not named_goods:
                        footer_reached = True
                        continue
                    if footer_reached:
                        if any(re.match(r'^'+CUSTOMER_LABEL+r'(?:\s|[:：]|$)',text) or re.match(r'^'+DATE_LABEL+r'(?:\s|[:：]|$)',text) for text in normalized):
                            later_info=metadata(sheet,row_index+1,row_index)
                            for field in ('customer','date'):
                                if later_info[field] and (not info[field] or normalize(later_info[field])!=normalize(info[field])):
                                    label='khách hàng' if field=='customer' else 'ngày xuất'
                                    issue('error',f'Dòng {row_index}: {label} sau dòng tổng khác hoặc chưa xác định ở bảng đầu. Hãy tách các phiếu ra sheet riêng.',row=row_index,field=field)
                        later_numbers={field:number(read(field)) for field in ('quantity','price','amount')}
                        goods_evidence=later_numbers['quantity'] is not None or later_numbers['price'] is not None or (later_numbers['amount'] is not None and bool(name))
                        if not goods_evidence:
                            if later_numbers['amount'] not in (None,Decimal(0)):
                                issue('warning',f'Dòng {row_index}: chỉ có số tiền sau dòng tổng, không có tên hàng/số lượng/đơn giá. Không đưa dòng tổng phụ này vào hàng; hãy kiểm tra.',row=row_index)
                                result['requires_review']=True
                                result['selected']=False
                            continue
                        issue('warning',f'Dòng {row_index}: có số liệu hàng sau dòng tổng. Tiếp tục đọc để tránh bỏ sót; hãy kiểm tra trước khi chọn sheet.',row=row_index)
                        footer_reached=False
                        result['requires_review']=True
                        result['selected']=False
                    if any(SIGNATURE.match(text) for text in normalized) and not any(number(sheet.cell(row_index,table_map[field]).value) is not None for field in ('quantity','price') if field in table_map): continue
                    if name is None or str(name).strip()=='':
                        quantity_raw,price_raw,amount_raw = read('quantity'),read('price'),read('amount')
                        # Empty rows containing unused =D9*E9 placeholders have no goods.
                        qty,price,amount = number(quantity_raw),number(price_raw),number(amount_raw)
                        if quantity_raw not in (None,'') or price_raw not in (None,'') or (amount is not None and amount!=0) or (amount is None and amount_raw not in (None,'')):
                            issue('error',f'Dòng {row_index}: có số liệu hàng nhưng thiếu tên hàng.',row=row_index,field='name')
                        continue
                    numeric_raw = {field:read(field) for field in ('quantity','price','amount')}
                    if field_of(name)=='name' and sum(field_of(value)==field for field,value in numeric_raw.items() if value is not None)>=2: continue
                    parsed = {field:number(value) for field,value in numeric_raw.items()}
                    qty,price,amount = parsed['quantity'],parsed['price'],parsed['amount']
                    # Named section rows with no quantity/price/amount are explanatory text.
                    if qty is None and price is None and amount is None:
                        populated = [value for field,value in numeric_raw.items() if value not in (None,'')]
                        if not populated:
                            issue('warning',f'Dòng {row_index}: bỏ qua dòng chữ không có số liệu hàng ({str(name)[:70]}).',row=row_index)
                            continue
                    if qty is None: raise ValueError('Thiếu hoặc không đọc được số lượng')
                    for field in ('price','amount'):
                        if numeric_raw[field] not in (None,'') and parsed[field] is None:
                            raise ValueError(f'Không đọc được {"đơn giá" if field=="price" else "thành tiền"}; cần sửa ô nguồn')
                    if price is None and amount is not None and qty!=0:
                        price = amount/qty
                        issue('warning',f'Dòng {row_index}: suy ra đơn giá từ thành tiền / số lượng.',row=row_index,field='price')
                    if price is None: raise ValueError('Thiếu hoặc không đọc được đơn giá')
                    if qty<0 or price<0 or (amount is not None and amount<0): raise ValueError('Có số âm; cần kiểm tra trước khi xuất kho')
                    if qty>Decimal('1e12') or price>Decimal('1e15'): raise ValueError('Số lượng hoặc đơn giá quá lớn; cần kiểm tra ô nguồn')
                    calculated = (qty*price).quantize(Decimal('1'),rounding=ROUND_HALF_UP)
                    if calculated>=10**18: raise ValueError('Thành tiền vượt phạm vi hỗ trợ')
                    if amount is None:
                        amount = calculated
                        issue('warning',f'Dòng {row_index}: tính thành tiền từ số lượng × đơn giá.',row=row_index,field='amount')
                    else:
                        amount = amount.quantize(Decimal('1'),rounding=ROUND_HALF_UP)
                        if amount>=10**18: raise ValueError('Thành tiền vượt phạm vi hỗ trợ')
                        if abs(amount-calculated)>1:
                            issue('error',f'Dòng {row_index}: thành tiền {amount} khác số lượng × đơn giá {calculated}. Cần sửa file nguồn hoặc dòng hàng.',row=row_index,field='amount',recoverable=True)
                    raw_qty = numeric_raw['quantity']
                    if isinstance(raw_qty,str) and re.fullmatch(r'[1-9]\d{0,2}[.,]\d{3}',raw_qty.strip()):
                        issue('warning',f'Dòng {row_index}: số lượng dạng chữ “{raw_qty}” có dấu phân cách chưa rõ. Đang đọc là {qty}; hãy kiểm tra và chỉnh số lượng.',row=row_index,field='quantity')
                        result['requires_review'] = True
                        result['selected'] = False
                    last_item_row = row_index
                    result['items'].append({'row':row_index,'name':str(name).strip(),'quantity':float(qty),'price':float(price),'amount':int(amount),
                                            'unit':str(read('unit') or ''),'code':str(read('code') or ''),'note':str(read('note') or '')})
                except Exception as exc:
                    issue('error',f'Dòng {row_index}: {exc}',row=row_index)
                    last_item_row = row_index
            previous_end = last_item_row+1
        result['total'] = sum(item['amount'] for item in result['items'])
        if not result['items']: issue('error','Không có dòng hàng hợp lệ.')
        if result['total']>=10**18: issue('error','Tổng tiền vượt phạm vi hỗ trợ.')
        formula_count = sum(1 for title,_ in values.evaluated if title==sheet.title)
        if formula_count: notes.append(f'Đã tính {formula_count} ô công thức chưa có kết quả lưu trong Excel.')
        if result['items']:
            result['confidence'] = min(99,result['confidence']+min(4,len(result['items'])))
        if errors: result['confidence'] = max(0,result['confidence']-min(35,10+5*len(errors)))
        if result['requires_review']: result['confidence'] = min(result['confidence'],65)
    return results


DIGITS=['không','một','hai','ba','bốn','năm','sáu','bảy','tám','chín']
def money_words(value):
    value=int(value)
    if value<0 or value>=10**18: raise ValueError('Tổng tiền ngoài phạm vi hỗ trợ')
    if value==0: return 'Không đồng.'
    def triple(n,full=False):
        h,t,u=n//100,n//10%10,n%10
        out=[]
        if h or full: out.extend([DIGITS[h],'trăm'])
        if t>1: out.extend([DIGITS[t],'mươi'])
        elif t==1: out.append('mười')
        elif u and (h or full): out.append('lẻ')
        if u: out.append('mốt' if u==1 and t>1 else 'lăm' if u==5 and t>=1 else DIGITS[u])
        return ' '.join(out)
    groups=[]
    while value: groups.append(value%1000); value//=1000
    units=['','nghìn','triệu','tỷ','nghìn tỷ','triệu tỷ']
    parts=[]
    for index in range(len(groups)-1,-1,-1):
        if groups[index]: parts.append((triple(groups[index],bool(parts) and groups[index]<100)+' '+units[index]).strip())
    text=' '.join(parts)+' đồng.'
    return text[0].upper()+text[1:]


def validate_template(data):
    wb=load_workbook(io.BytesIO(data))
    sheet=wb.active
    if normalize(sheet['A5'].value)!='phieu xuat kho' or field_of(sheet['B11'].value)!='name' or 'tong cong' not in normalize(sheet['B27'].value):
        raise ValueError('Template phải theo mẫu đã cung cấp: tiêu đề A5, bảng hàng dòng 11 và tổng cộng dòng 27.')
    return wb


def put(sheet,coord,value):
    sheet[coord]=value
    if isinstance(value,str): sheet[coord].data_type='s'


def wrap_cell(sheet, coord, width=None):
    cell=sheet[coord]
    if cell.value is None: return 1
    alignment=copy.copy(cell.alignment)
    alignment.wrap_text=True
    alignment.vertical='center'
    cell.alignment=alignment
    available=width or max(4,sheet.column_dimensions[cell.column_letter].width-2)
    return max(1,sum(max(1,math.ceil(len(line)/available)) for line in str(cell.value).split('\n')))


def export_items(items, exclude_zero=False):
    validated=[]
    for index,item in enumerate(items,1):
        name=str(item.get('name','')).strip()
        qty,price,amount=(number(item.get(key)) for key in ('quantity','price','amount'))
        if not name or qty is None or price is None or amount is None:
            raise ValueError(f'Dòng hàng {index}: thiếu tên hàng hoặc số liệu không hợp lệ.')
        if min(qty,price,amount)<0 or qty>Decimal('1e12') or price>Decimal('1e15') or amount>=10**18:
            raise ValueError(f'Dòng hàng {index}: số liệu ngoài phạm vi hỗ trợ.')
        calculated=(qty*price).quantize(Decimal('1'),rounding=ROUND_HALF_UP)
        amount=amount.quantize(Decimal('1'),rounding=ROUND_HALF_UP)
        if abs(amount-calculated)>1:
            raise ValueError(f'Dòng hàng {index}: thành tiền khác số lượng × đơn giá.')
        if exclude_zero and qty==0: continue
        validated.append({**item,'name':name,'quantity':float(qty),'price':float(price),'amount':int(amount)})
    return validated


def export(template, sheets, exclude_zero=False):
    wb=validate_template(template)
    base=wb.active
    selected=[s for s in sheets if s.get('selected')]
    if not selected: raise ValueError('Chọn ít nhất một sheet để xuất.')
    used=set()
    totals={}
    for entry in selected:
        if entry.get('errors'): raise ValueError(f"Sheet {entry['sheet']} có lỗi dữ liệu. Hãy sửa trước khi xuất.")
        customer=str(entry.get('customer','')).strip()
        if not customer: raise ValueError(f"Chưa nhập khách hàng cho sheet {entry['sheet']}.")
        try: day=datetime.strptime(entry.get('date',''),'%Y-%m-%d')
        except (ValueError,TypeError): raise ValueError(f"Ngày xuất không hợp lệ: {entry['sheet']}")
        items=export_items(entry['items'],exclude_zero)
        if not items: raise ValueError(f"Sheet {entry['sheet']} không còn dòng hàng để xuất.")
        sheet=wb.copy_worksheet(base)
        title=re.sub(r'[\\/*?:\[\]]','-',entry['sheet'])[:31] or 'Phiếu xuất'
        candidate=title; index=2
        while candidate.lower() in used or candidate==base.title:
            candidate=title[:26]+f' ({index})'; index+=1
        used.add(candidate.lower()); sheet.title=candidate
        extra=max(0,len(items)-15)
        if extra:
            shifted=[]
            for merged in list(sheet.merged_cells.ranges):
                if merged.min_row>=27:
                    shifted.append((merged.min_col,merged.min_row+extra,merged.max_col,merged.max_row+extra))
                    sheet.unmerge_cells(str(merged))
            dims={r:copy.copy(d) for r,d in sheet.row_dimensions.items() if r>=27}
            sheet.insert_rows(27,extra)
            for old,dim in sorted(dims.items(),reverse=True):
                if old in sheet.row_dimensions: del sheet.row_dimensions[old]
                dim.index=old+extra; sheet.row_dimensions[old+extra]=dim
            for bounds in shifted: sheet.merge_cells(start_column=bounds[0],start_row=bounds[1],end_column=bounds[2],end_row=bounds[3])
        total_row=27+extra
        for r in range(12,total_row):
            for c in range(1,10):
                cell=sheet.cell(r,c)
                if r>=27: cell._style=copy.copy(sheet.cell(12,c)._style)
                cell.value=None
            if r>=27: sheet.row_dimensions[r].height=sheet.row_dimensions[12].height
        put(sheet,'B6',entry.get('voucher',''))
        sheet['E6']=day
        put(sheet,'B7',customer)
        put(sheet,'B8',entry.get('recipient',''))
        put(sheet,'B9',entry.get('address',''))
        if entry.get('content'): put(sheet,'B10',entry['content'])
        for coord,width in (('B7',85),('B8',85),('B9',120),('B10',120)):
            lines=wrap_cell(sheet,coord,width)
            row=sheet[coord].row
            sheet.row_dimensions[row].height=max(sheet.row_dimensions[row].height or 21,lines*15)
        for index,item in enumerate(items,1):
            r=index+11
            vals=[index,item['name'],item.get('code',''),item.get('unit',''),item['quantity'],item['quantity'],item['price'],item['amount'],item.get('note','')]
            for c,value in enumerate(vals,1): put(sheet,f'{get_column_letter(c)}{r}',value)
            lines=max(wrap_cell(sheet,f'{col}{r}') for col in ('B','C','D','I'))
            sheet.row_dimensions[r].height=max(sheet.row_dimensions[r].height or 23.1,lines*14)
        total=sum(item['amount'] for item in items)
        # Numeric amounts preserve source totals and display in readers without formula recalculation.
        sheet.cell(total_row,8,total)
        put(sheet,f'B{29+extra}',money_words(total))
        sheet.print_area=f'A1:I{32+extra}'
        sheet.print_options.horizontalCentered=True
        if not sheet.page_setup.paperSize: sheet.page_setup.paperSize=sheet.PAPERSIZE_A4
        if not sheet.page_setup.orientation: sheet.page_setup.orientation=sheet.ORIENTATION_PORTRAIT
        sheet.sheet_properties.pageSetUpPr.fitToPage=True
        sheet.page_setup.fitToWidth=1; sheet.page_setup.fitToHeight=0
        sheet.print_title_rows='1:11'
        totals[sheet.title]=total
    wb.remove(base)
    # The supplied template contains one worksheet; remove any extra template-only sheets.
    for sheet in list(wb):
        if sheet.title not in totals: wb.remove(sheet)
    wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
    result=io.BytesIO(); wb.save(result)
    return result.getvalue()
