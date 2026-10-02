"""Convert browser OCR words or a reviewed text table into an auditable Excel source.

OCR is deliberately a proposal: raw numeric strings are preserved, every uncertain
row stays visible, and the HTTP caller must require confirmation before exporting.
No external OCR service, secrets, or optional native dependencies are used here.
"""
from __future__ import annotations

import csv
import io
import math
import re
from datetime import date
from difflib import SequenceMatcher
from statistics import median
from pathlib import PurePath

from dependencies import configure_vendor

configure_vendor()

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from engine import ALIASES, normalize, number

MAX_DOCUMENTS = 10
MAX_WORDS_PER_DOCUMENT = 50000
MAX_TOTAL_WORDS = 200000
MAX_TSV_BYTES = 8_000_000
MAX_TEXT_CHARS = 1_000_000
FIELDS = ('name', 'code', 'unit', 'quantity', 'price', 'amount', 'note')
LABELS = {'name': 'Tên hàng', 'code': 'Mã hàng', 'unit': 'ĐVT', 'quantity': 'Số lượng',
          'price': 'Đơn giá', 'amount': 'Thành tiền', 'note': 'Ghi chú'}
METADATA_LABELS = {'customer': 'Khách hàng', 'date': 'Ngày xuất', 'address': 'Địa chỉ',
                   'voucher': 'Số phiếu', 'recipient': 'Người nhận hàng'}
META_PATTERN = re.compile(
    r'\b(ten khach hang|khach hang|ho ten nguoi mua hang|customer(?: name)?|buyer|don vi mua(?: hang)?|'
    r'ngay xuat(?: kho)?|ngay giao(?: hang)?|ngay lap|ngay|date|'
    r'dia chi(?: giao hang)?|address|so phieu(?: xuat(?: kho)?)?|voucher(?: no)?|'
    r'nguoi nhan(?: hang)?|recipient)\s*[:：]?\s*', re.I)
FOOTER_PATTERN = re.compile(r'^(?:tong cong|tong tien(?: hang| thanh toan)?|tong thanh tien|cong tien(?: hang)?|cong thanh tien|grand total|subtotal|total(?: amount)?|cong|tong)\s*[:：]?$', re.I)
SIGNATURE_PATTERN = re.compile(r'^(?:nguoi lap(?: phieu)?|thu kho|ke toan|giam doc|ky(?: ten)?|signature|prepared by|approved by|bang chu|so tien bang chu)(?:\b|:)', re.I)
UNITS = {'kg','g','gam','tan','ta','cai','chiec','hop','chai','tui','thung','mo','bo','suat','phan','khay','qua','con','lit','l','ml','pcs','can','goi','tap'}


def _issue(report, level, message, row=None, field=None, source_row=None):
    issue = {'level': level, 'message': message, 'recoverable': False}
    if row is not None: issue['row'] = row
    if source_row is not None: issue['source_row'] = source_row
    if field: issue['field'] = field
    report['issues'].append(issue)
    if level == 'warning': report['warnings'].append(message)


def _normalized_header(text):
    text = normalize(text)
    text = re.sub(r'\([^)]*\)|\[[^]]*\]', '', text)
    text = re.sub(r'[.,:_/\-]+', ' ', text)
    return re.sub(r'\s+', ' ', text).strip()


HEADER_ALIASES = {key: {_normalized_header(alias) for alias in values} for key, values in ALIASES.items()}
HEADER_ALIASES['name'].update({'ten tp', 'ten thuc pham', 'thuc pham'})
HEADER_ALIASES['stt'] = {'stt', 's t t', 'so tt', 'no', 'number', '#', 'tt'}


def _header_label(text):
    clean = _normalized_header(text)
    compact = clean.replace(' ', '')
    if compact in ('dvt', 'sl', 'dg', 'qty', 'uom', 'sku', 'stt'):
        clean = compact
    if not clean: return None
    best = None
    for field, aliases in HEADER_ALIASES.items():
        for alias in aliases:
            score = 1.0 if clean == alias or (len(alias.replace(' ', '')) >= 4 and compact == alias.replace(' ', '')) else 0.0
            if not score and clean.startswith(alias+' '):
                suffix = clean[len(alias):].strip()
                if re.fullmatch(r'(?:vnd|dong|kg|gam|g|lit|l|vn d|d|usd)(?:\s+(?:kg|g|lit|l))?', suffix): score = .96
            if not score and len(clean) >= 5 and len(alias) >= 5 and .70 <= len(clean)/len(alias) <= 1.35:
                ratio = SequenceMatcher(None,clean,alias).ratio()
                if ratio >= .84: score = ratio*.93
            if score and (best is None or score > best[1] or (score == best[1] and len(alias)>best[2])):
                best = (field,score,len(alias))
    return (best[0],best[1]) if best else None


def _metadata(text):
    # Keep original Vietnamese text: accents change matching, not customer names.
    chars, positions = [], []
    import unicodedata
    for index,char in enumerate(text):
        decomposed = unicodedata.normalize('NFD',char.lower().replace('đ','d'))
        for part in decomposed:
            if not unicodedata.combining(part): chars.append(part);positions.append(index)
    normalized = ''.join(chars)
    matches = []
    for match in META_PATTERN.finditer(normalized):
        fragment=normalized[match.start():match.end()]
        if match.start()>0 and ':' not in fragment and '：' not in fragment:continue
        if ':' not in fragment and '：' not in fragment:
            trailing=normalized[match.end():].strip()
            if not match[1].startswith(('ngay','date')) or not re.match(r'^(?:\d{1,2}[/.\-]\d{1,2}[/.\-]\d{4}|\d{4}[/.\-]\d{1,2}[/.\-]\d{1,2}|\d{1,2}\s+thang\s+\d{1,2}\s+nam\s+\d{4})',trailing):continue
        matches.append(match)
    result = {}
    for index,match in enumerate(matches):
        label = match[1]
        key = 'customer' if any(label.startswith(token) for token in ('ten khach','khach','ho ten','customer','buyer','don vi')) else 'date' if label.startswith(('ngay','date')) else 'address' if label.startswith(('dia chi','address')) else 'voucher' if label.startswith(('so phieu','voucher')) else 'recipient'
        start = positions[match.end()] if match.end()<len(positions) else len(text)
        stop_norm = matches[index+1].start() if index+1<len(matches) else len(normalized)
        stop = positions[stop_norm] if stop_norm<len(positions) else len(text)
        value = text[start:stop].strip(' :：|\t')
        if key == 'date':
            original = normalize(value)
            full = re.search(r'(\d{1,2})\s+thang\s+(\d{1,2})\s+nam\s+(\d{4})', original)
            normal = re.search(r'\b(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{4})\b', original)
            iso = re.search(r'\b(\d{4})[/.\-](\d{1,2})[/.\-](\d{1,2})\b', original)
            try:
                if full or normal:
                    parsed = full or normal
                    value = date(int(parsed[3]),int(parsed[2]),int(parsed[1])).isoformat()
                elif iso: value = date(int(iso[1]),int(iso[2]),int(iso[3])).isoformat()
            except ValueError: pass
        if value: result[key] = value
    return result


def _parse_words(document, report):
    tsv = document.get('tsv','')
    if not isinstance(tsv,str) or len(tsv.encode('utf-8')) > MAX_TSV_BYTES: raise ValueError('Dữ liệu OCR TSV quá lớn hoặc không hợp lệ.')
    if not tsv.strip(): return []
    first_line=tsv.lstrip('\ufeff').splitlines()[0].split('\t')
    fieldnames=['level','page_num','block_num','par_num','line_num','word_num','left','top','width','height','conf','text'] if len(first_line)>=12 and first_line[0] in ('1','2','3','4','5') else None
    reader = csv.DictReader(io.StringIO(tsv.lstrip('\ufeff')),delimiter='\t',fieldnames=fieldnames)
    required = {'level','left','top','width','height','conf','text'}
    if not reader.fieldnames or not required.issubset(reader.fieldnames): raise ValueError('OCR TSV thiếu cột vị trí, độ tin cậy hoặc nội dung.')
    words = []
    count = 0
    for record in reader:
        if record.get('level') != '5': continue
        count += 1
        if count>MAX_WORDS_PER_DOCUMENT: raise ValueError('Mỗi ảnh được hỗ trợ tối đa 50.000 từ OCR.')
        text = (record.get('text') or '').strip()
        if not text: continue
        try:
            x,y,w,h = (float(record[key]) for key in ('left','top','width','height'))
            confidence = float(record['conf'])
            if not all(math.isfinite(value) for value in (x,y,w,h,confidence)) or min(x,y)<0 or min(w,h)<=0 or max(x+w,y+h)>100000 or not 0<=confidence<=100:
                raise ValueError()
        except (ValueError,TypeError) as exc: raise ValueError('Từ OCR có vị trí hoặc độ tin cậy không hợp lệ. Hãy nhận diện lại hoặc sửa bảng chữ.') from exc
        if len(text)>1000:raise ValueError('Một từ OCR quá dài; hãy nhận diện lại hoặc sửa bảng chữ.')
        words.append({'text':text,'x':x,'y':y,'w':w,'h':h,'cx':x+w/2,'cy':y+h/2,'right':x+w,'bottom':y+h,'conf':confidence})
    return words


def _visual_rows(words):
    rows = []
    for word in sorted(words,key=lambda word:(word['cy'],word['x'])):
        candidates = []
        for index in range(max(0,len(rows)-5),len(rows)):
            row = rows[index]
            overlap = min(row['bottom'],word['bottom'])-max(row['top'],word['y'])
            if overlap/max(1,min(row['height'],word['h'])) >= .38 or abs(row['cy']-word['cy']) <= max(4,min(row['height'],word['h'])*.42):
                candidates.append((abs(row['cy']-word['cy']),index))
        if candidates:
            row = rows[min(candidates)[1]]
            row['words'].append(word)
            row['top']=min(row['top'],word['y']);row['bottom']=max(row['bottom'],word['bottom'])
            row['height']=median([item['h'] for item in row['words']]);row['cy']=median([item['cy'] for item in row['words']])
        else:
            rows.append({'words':[word],'top':word['y'],'bottom':word['bottom'],'height':word['h'],'cy':word['cy']})
    rows.sort(key=lambda row:row['cy'])
    for index,row in enumerate(rows,1):
        row['words'].sort(key=lambda word:word['x'])
        row['source_row']=index
        row['text']=' '.join(word['text'] for word in row['words'])
        row['confidence']=round(sum(word['conf']*max(1,len(word['text'])) for word in row['words'])/sum(max(1,len(word['text'])) for word in row['words']),1)
        row['bbox']={'left':min(word['x'] for word in row['words']),'top':row['top'],'width':max(word['right'] for word in row['words'])-min(word['x'] for word in row['words']),'height':row['bottom']-row['top']}
    return rows


def _header_fragments(row):
    words = row['words'];fragments=[]
    for start in range(len(words)):
        for size in range(1,min(6,len(words)-start)+1):
            group=words[start:start+size]
            if any(group[index+1]['x']-group[index]['right']>max(35,row['height']*2.5) for index in range(len(group)-1)): break
            match=_header_label(' '.join(word['text'] for word in group))
            if match:
                fragments.append({'field':match[0],'score':match[1],'x':group[0]['x'],'right':max(word['right'] for word in group),'cx':(group[0]['x']+max(word['right'] for word in group))/2,'text':' '.join(word['text'] for word in group),'words':group})
    return fragments


def _header_at(rows,index):
    first=rows[index]
    possibilities=[]
    for span in (1,2):
        fragments=_header_fragments(first)
        if span==2:
            if index+1>=len(rows) or rows[index+1]['top']-first['bottom']>first['height']*1.7:continue
            if _metadata(first['text']) or any(number(word['text']) is not None for word in first['words']):continue
            second=rows[index+1]
            if any(number(word['text']) is not None for word in second['words']): continue
            fragments+=_header_fragments(second)
            for upper in first['words']:
                for lower in second['words']:
                    if abs(upper['cx']-lower['cx']) <= max(upper['w'],lower['w'])*.75+8:
                        match=_header_label(upper['text']+' '+lower['text'])
                        if match:
                            fragments.append({'field':match[0],'score':match[1],'x':min(upper['x'],lower['x']),'right':max(upper['right'],lower['right']),'cx':(upper['cx']+lower['cx'])/2,'text':upper['text']+' '+lower['text'],'words':[upper,lower]})
        mapping={}
        for fragment in sorted(fragments,key=lambda f:(f['score'],len(f['text'])),reverse=True):
            field=fragment['field']
            if field not in mapping and not any(abs(fragment['cx']-existing['cx'])<max(10,first['height']*.45) for existing in mapping.values()):mapping[field]=fragment
        if 'name' in mapping and 'quantity' in mapping and ('price' in mapping or 'amount' in mapping):
            if span==2 and not any(any(word is upper for upper in first['words'])
                                  for anchor in mapping.values() for word in anchor['words']):
                continue
            confidence=round(sum(anchor['score'] for anchor in mapping.values())/len(mapping)*100)
            possibilities.append({'index':index,'span':span,'anchors':mapping,'confidence':confidence})
    return max(possibilities,key=lambda p:(len(p['anchors']),p['confidence'],-p['span'])) if possibilities else None


def _numeric_clusters(row, anchors=None):
    clusters=[]
    name_bounds=None
    if anchors and 'name' in anchors:
        ordered=sorted(anchors.items(),key=lambda pair:pair[1]['cx'])
        for index,(field,anchor) in enumerate(ordered):
            if field=='name':
                left=-math.inf if index==0 else (ordered[index-1][1]['cx']+anchor['cx'])/2
                right=math.inf if index==len(ordered)-1 else (ordered[index+1][1]['cx']+anchor['cx'])/2
                name_bounds=(left,right)
    for word in row['words']:
        text=word['text']
        numericish=bool(re.search(r'\d',text)) and sum(char.isdigit() or char in '.,+-()₫đ' for char in text)/max(1,len(text))>=.55
        if not numericish: continue
        if name_bounds and name_bounds[0]<=word['cx']<name_bounds[1]:continue
        if clusters:
            previous=clusters[-1]
            join_gap=word['x']-previous['right']
            if join_gap<=max(8,row['height']*.5) and (re.fullmatch(r'\d{3}',text) or previous['text'].endswith(('.',','))):
                previous['text']+=(' ' if not previous['text'].endswith(('.',',')) else '')+text
                previous['words'].append(word);previous['right']=word['right'];previous['cx']=(previous['x']+previous['right'])/2
                continue
        clusters.append({'text':text,'words':[word],'x':word['x'],'right':word['right'],'cx':word['cx']})
    return clusters


def _assign_numeric(clusters, anchors):
    fields=sorted((field for field in anchors if field in ('stt','code','quantity','price','amount')),key=lambda field:anchors[field]['cx'])
    if not fields or not clusters:return {},set()
    distances=[abs(anchors[fields[index+1]]['cx']-anchors[fields[index]]['cx']) for index in range(len(fields)-1)]
    scale=max(30,median(distances) if distances else 100)
    # Ordered dynamic assignment permits missing cells while respecting right alignment.
    memo={}
    def solve(i,j):
        if i==len(clusters):return (0,[])
        if j==len(fields):return ((len(clusters)-i)*2.8,[])
        key=(i,j)
        if key in memo:return memo[key]
        options=[]
        cost,path=solve(i,j+1);options.append((cost+.75,path))
        cost,path=solve(i+1,j);options.append((cost+2.8,path))
        anchor=anchors[fields[j]]
        distance=abs(clusters[i]['right']-anchor.get('numeric_right',anchor['right']))/scale
        cost,path=solve(i+1,j+1);options.append((cost+min(5,distance),[(i,fields[j])]+path))
        result=min(options,key=lambda option:option[0]);memo[key]=result;return result
    _,pairs=solve(0,0)
    return {field:clusters[index] for index,field in pairs},{id(word) for index,_ in pairs for word in clusters[index]['words']}


def _row_cells(row, anchors):
    ordered=sorted(anchors.items(),key=lambda pair:pair[1]['cx'])
    groups={field:[] for field in FIELDS}
    numeric,used=_assign_numeric(_numeric_clusters(row,anchors),anchors)
    for field,cluster in numeric.items():
        if field in groups:groups[field].extend(cluster['words'])
    for word in row['words']:
        if id(word) in used:continue
        chosen=min(ordered,key=lambda pair:abs(word['cx']-pair[1]['cx']))[0]
        for index,(field,anchor) in enumerate(ordered):
            left=-math.inf if index==0 else (ordered[index-1][1]['cx']+anchor['cx'])/2
            right=math.inf if index==len(ordered)-1 else (ordered[index+1][1]['cx']+anchor['cx'])/2
            if left<=word['cx']<right:chosen=field;break
        if chosen in groups:groups[chosen].append(word)
    cells={field:' '.join(word['text'] for word in sorted(group,key=lambda word:word['x'])) for field,group in groups.items()}
    confidence={field:round(sum(word['conf']*max(1,len(word['text'])) for word in group)/sum(max(1,len(word['text'])) for word in group),1) for field,group in groups.items() if group}
    return cells,confidence,numeric


def _fallback_rows(rows):
    candidates=[]
    for row in rows:
        numeric=_numeric_clusters(row)
        # An STT number precedes the text; remaining three numbers form a possible item.
        textwords=[word for word in row['words'] if re.search(r'[A-Za-zÀ-ỹ]',word['text']) and not re.fullmatch(r'[\d.,]+',word['text'])]
        if not textwords:continue
        first_text=min(word['x'] for word in textwords)
        numeric=[cluster for cluster in numeric if cluster['x']>=first_text]
        if len(numeric)!=3:continue
        qty,price,amount=(number(cluster['text']) for cluster in numeric)
        if qty is None or price is None or amount is None or not 0<qty<price or price<100 or abs(qty*price-amount)>1:continue
        prefix=[word for word in row['words'] if word['right']<=numeric[0]['x'] and word['x']>=first_text]
        unit=''
        if prefix and normalize(prefix[-1]['text']) in UNITS:unit=prefix.pop()['text']
        name=' '.join(word['text'] for word in prefix)
        if not name:continue
        cells={field:'' for field in FIELDS};cells.update(name=name,unit=unit,quantity=numeric[0]['text'],price=numeric[1]['text'],amount=numeric[2]['text'])
        candidates.append({'rowdata':row,'cells':cells,'cell_confidence':{}})
    if len(candidates)<2:return None
    # Consistency across two rows is evidence, never automatic approval.
    anchors={}
    for field,index in (('quantity',0),('price',1),('amount',2)):
        matching=[_numeric_clusters(candidate['rowdata'])[-3+index] for candidate in candidates]
        anchors[field]={'cx':median(cluster['cx'] for cluster in matching),'right':median(cluster['right'] for cluster in matching),'numeric_right':median(cluster['right'] for cluster in matching)}
    anchors['name']={'cx':median(min(word['x'] for word in candidate['rowdata']['words'] if re.search(r'[A-Za-zÀ-ỹ]',word['text'])) for candidate in candidates),'right':0}
    return anchors


def _plain_rows(text, edited=False):
    rows=[]
    for index,line in enumerate(text.splitlines(),1):
        if not line.strip():continue
        separator = '\t' if '\t' in line else '|' if '|' in line else None
        split = next(csv.reader([line], delimiter=separator, skipinitialspace=True)) if separator else re.split(r'\s{2,}', line.strip())
        parts=[part.strip() for part in split]
        rows.append({'source_row':index,'text':line.strip(),'parts':parts,'confidence':None if edited else 0,'bbox':None})
    return rows


def _plain_extract(rows, report):
    headers=None;header_index=-1
    for index,row in enumerate(rows):
        detected={}
        for col,part in enumerate(row['parts']):
            label=_header_label(part)
            if label and label[0] not in detected:detected[label[0]]=col
        if 'name' in detected and 'quantity' in detected and ('price' in detected or 'amount' in detected):headers=detected;header_index=index;break
    if headers is None:
        # The review editor's documented layout is name | unit | qty | price | amount.
        eligible=[(index,row) for index,row in enumerate(rows) if len(row['parts'])>=3 and not _metadata(row['text'])]
        if eligible:
            header_index=eligible[0][0]-1
            count=len(eligible[0][1]['parts'])
            headers={'name':0,'unit':1,'quantity':2,'price':3,'amount':4} if count>=5 else {'name':0,'quantity':1,'price':2,**({'amount':3} if count>=4 else {})}
            _issue(report,'warning','Bảng chữ chưa có tiêu đề; đang dùng thứ tự cột trong hướng dẫn. Hãy kiểm tra từng dòng.')
    entries=[];info={}
    for index,row in enumerate(rows):
        structured_goods=len(row.get('parts',[]))>=3 and any(number(part) is not None for part in row.get('parts',[])[1:])
        meta={} if structured_goods else _metadata(row['text'])
        if meta:
            if entries:entries.append({'rowdata':row,'metadata':meta})
            else:info.update(meta)
            continue
        if index==header_index:continue
        if index<header_index and len(row['parts'])<2:continue
        if headers is None:
            if SIGNATURE_PATTERN.match(normalize(row['text'])):continue
            cells={field:'' for field in FIELDS};cells['name']=row['text'];cells['quantity']='[OCR: chưa rõ cột số lượng]'
            entries.append({'rowdata':row,'cells':cells,'cell_confidence':{}});continue
        parts=row['parts']
        new_headers={}
        for col,part in enumerate(parts):
            label=_header_label(part)
            if label:new_headers[label[0]]=col
        if 'name' in new_headers and 'quantity' in new_headers and ('price' in new_headers or 'amount' in new_headers):
            headers=new_headers
            continue
        first=normalize(parts[headers['name']] if headers['name']<len(parts) else '')
        if FOOTER_PATTERN.fullmatch(first) or SIGNATURE_PATTERN.match(first):
            report['source_rows'].append({'source_row':row['source_row'],'text':row['text'],'confidence':row['confidence'],'type':'footer'})
            continue
        cells={field:(parts[col] if col<len(parts) else '') for field,col in headers.items() if field in FIELDS}
        cells={field:cells.get(field,'') for field in FIELDS}
        # Unsplit malformed rows still have a name, and a marker prevents silent omission.
        if not cells['name'] and not any(cells[field] for field in ('quantity','price','amount')):
            cells['name']=row['text'];cells['quantity']='[OCR: không tách được dòng]'
        entries.append({'rowdata':row,'cells':cells,'cell_confidence':{}})
    if headers is None:_issue(report,'error','Chưa nhận diện được các cột trong bảng chữ. Hãy sửa theo mẫu Tên hàng | ĐVT | Số lượng | Đơn giá | Thành tiền.')
    report['detection_method']='edited-text' if report.get('edited_text') else 'ocr-text'
    return entries,info


def _tsv_extract(rows, report):
    best=None
    for index in range(len(rows)):
        candidate=_header_at(rows,index)
        if candidate:
            best=candidate
            break
    info={};prefix_entries=[]
    if best:
        header_end=best['index']+best['span'];anchors=best['anchors']
        report['detection_method']='ocr-headers';report['header_confidence']=best['confidence']
        for row in rows[:best['index']]:
            meta=_metadata(row['text'])
            if meta:
                if prefix_entries:prefix_entries.append({'rowdata':row,'metadata':meta})
                else:info.update(meta)
                continue
            if len(_numeric_clusters(row,anchors))>=2:
                cells,cell_confidence,_=_row_cells(row,anchors)
                prefix_entries.append({'rowdata':row,'cells':cells,'cell_confidence':cell_confidence})
                _issue(report,'warning','Có dòng có số liệu trước tiêu đề rõ ràng đầu tiên; giữ lại để đối chiếu thay vì bỏ qua.',source_row=row['source_row'])
        customer_candidates=[row['text'] for row in rows[:best['index']] if re.match(r'^(?:truong|cong ty|nha hang|khach san|mam non)\b',normalize(row['text']))]
        if not info.get('customer') and len(customer_candidates)==1:
            info['customer']=customer_candidates[0]
            _issue(report,'warning','Tên khách hàng được đề xuất từ dòng tiêu đề của ảnh. Hãy đối chiếu với bên mua thực tế.',field='customer')
        # Learn right-aligned numeric positions from complete rows below the header.
        right_edges={field:[] for field in ('quantity','price','amount')}
        for row_index,row in enumerate(rows[header_end:header_end+30],header_end):
            if _header_at(rows,row_index) or _metadata(row['text']):break
            cells,_,numeric=_row_cells(row,anchors)
            qty,price,amount=(number(cells[field]) for field in ('quantity','price','amount'))
            if qty is not None and price is not None and amount is not None and abs(qty*price-amount)<=1:
                for field in right_edges:
                    if field in numeric:right_edges[field].append(numeric[field]['right'])
        for field,edges in right_edges.items():
            if edges and field in anchors:anchors[field]['numeric_right']=median(edges)
    else:
        anchors=_fallback_rows(rows)
        header_end=0
        report['detection_method']='ocr-inferred' if anchors else 'none'
        if anchors:_issue(report,'warning','Không tìm được tiêu đề rõ ràng. Đề xuất cột từ ít nhất hai dòng có số lượng × đơn giá khớp thành tiền; hãy kiểm tra.')
        else:_issue(report,'error','Chưa tìm được tiêu đề hoặc đủ bằng chứng xác định cột từ ảnh. Hãy sửa bảng chữ trước khi nhận diện lại.')
    entries=prefix_entries
    skip_until=header_end
    for index,row in enumerate(rows[header_end:],header_end):
        if index<skip_until:continue
        repeated=_header_at(rows,index)
        if repeated and index>=header_end:
            anchors=repeated['anchors']
            skip_until=index+repeated['span']
            continue
        pre_cells=_row_cells(row,anchors)[0] if anchors else {}
        table_numbers=any(pre_cells.get(field) for field in ('quantity','price','amount'))
        explicit_label=bool(re.match(r'^(?:khach hang|ten khach hang|customer(?: name)?|ngay(?: xuat)?|date|dia chi|address|so phieu|voucher|nguoi nhan|recipient)\s*[:：]',normalize(row['text'])))
        valid_numbers=sum(number(pre_cells.get(field)) is not None for field in ('quantity','price','amount'))
        meta=_metadata(row['text']) if not table_numbers or (explicit_label and valid_numbers<2) else {}
        if meta:
            if entries:entries.append({'rowdata':row,'metadata':meta})
            else:info.update(meta)
            continue
        if anchors is None:
            if SIGNATURE_PATTERN.match(normalize(row['text'])):continue
            cells={field:'' for field in FIELDS};cells.update(name=row['text'],quantity='[OCR: chưa rõ cột số lượng]')
            entries.append({'rowdata':row,'cells':cells,'cell_confidence':{}})
            continue
        cells,cell_confidence,_=_row_cells(row,anchors)
        if FOOTER_PATTERN.fullmatch(normalize(cells['name'])):
            report['source_rows'].append({'source_row':row['source_row'],'text':row['text'],'confidence':row['confidence'],'bbox':row['bbox'],'type':'footer'})
            continue
        if SIGNATURE_PATTERN.match(normalize(cells['name'])) and not any(number(cells[field]) is not None for field in ('quantity','price','amount')):
            continue
        if not any(cells[field] for field in ('quantity','price','amount')) and entries and 'cells' in entries[-1]:
            previous=entries[-1]
            prior=previous['rowdata']
            if row['cy']-prior.get('cy',-1000)<=row['height']*1.5 and cells['name'] and not any(cells[field] for field in ('unit','code')):
                previous['cells']['name']+=' '+cells['name']
                previous.setdefault('continuation_rows',[]).append(row['source_row'])
                continue
        entries.append({'rowdata':row,'cells':cells,'cell_confidence':cell_confidence})
    return entries,info


def _sheet_title(filename, used):
    title=re.sub(r'[\\/*?:\[\]]','-',PurePath(filename).stem).strip()[:31] or 'Ảnh'
    candidate=title;index=2
    while candidate.casefold() in used:
        suffix=f' ({index})';candidate=title[:31-len(suffix)]+suffix;index+=1
    used.add(candidate.casefold())
    return candidate


def build_image_workbook(documents):
    """Return ``(xlsx_bytes, report)`` from Tesseract TSV or corrected text documents.

    Report source rows use the resulting workbook's ``row`` and the visual OCR
    ``source_row``. Numeric strings remain literal strings, including OCR errors.
    Callers should run engine.analyze and require a user review for every image.
    """
    if not isinstance(documents,list) or not 1<=len(documents)<=MAX_DOCUMENTS:
        raise ValueError('Chọn từ 1 đến 10 ảnh cho mỗi lần nhận diện.')
    from image_tables import expand_customers
    documents, customer_info = expand_customers(documents)
    wb=Workbook();wb.remove(wb.active)
    report={'documents':[],'image_info':{},'warnings':[]}
    used=set();total_words=0
    for index,document in enumerate(documents,1):
        if not isinstance(document,dict):raise ValueError('Thông tin ảnh OCR không hợp lệ.')
        filename=document.get('filename') or f'Ảnh {index}'
        if not isinstance(filename,str) or len(filename)>250:raise ValueError('Tên ảnh không hợp lệ hoặc quá dài.')
        text=document.get('text','')
        if not isinstance(text,str) or len(text)>MAX_TEXT_CHARS:raise ValueError('Nội dung OCR quá lớn hoặc không hợp lệ.')
        title=_sheet_title(filename,used);sheet=wb.create_sheet(title)
        docreport={'sheet':title,'image_name':filename,'filename':filename,'ocr_text':text,'ocr_confidence':0,'confidence':0,
                   'needs_review':True,'rotation':document.get('rotation',0),'language':str(document.get('language','vie+eng'))[:50],
                   'edited_text':bool(document.get('edited_text')),'confidence_origin':'edited-text' if document.get('edited_text') else 'ocr','warnings':[],'issues':[],'source_rows':[]}
        words=_parse_words(document,docreport) if not document.get('edited_text') else []
        total_words+=len(words)
        if total_words>MAX_TOTAL_WORDS:raise ValueError('Tổng nội dung OCR quá lớn; hãy chia thành nhiều lần nhập.')
        if words:
            visual=_visual_rows(words)
            entries,metadata=_tsv_extract(visual,docreport)
            docreport['ocr_confidence']=round(sum(word['conf']*max(1,len(word['text'])) for word in words)/sum(max(1,len(word['text'])) for word in words),1)
            if not text:docreport['ocr_text']='\n'.join(row['text'] for row in visual)
        else:
            visual=_plain_rows(text,bool(document.get('edited_text')))
            entries,metadata=_plain_extract(visual,docreport)
            docreport['ocr_confidence']=None if document.get('edited_text') else 0
            if not document.get('edited_text'):_issue(docreport,'warning','Ảnh chưa có vị trí từ OCR; đang đọc bảng chữ. Hãy kiểm tra việc tách cột.')
        docreport['confidence']=docreport['ocr_confidence']
        docreport['metadata']=dict(metadata)
        _issue(docreport,'warning','Dữ liệu từ ảnh cần đối chiếu tên hàng, số lượng, đơn giá, thành tiền và khách hàng trước khi chọn xuất.')
        for field,label in METADATA_LABELS.items():
            value=metadata.get(field)
            if value:sheet.append([label+': '+value])
        header=[LABELS[field] for field in FIELDS]
        sheet.append(header)
        header_row=sheet.max_row
        goods_count=0
        for entry in entries:
            rowdata=entry['rowdata']
            if 'metadata' in entry:
                docreport['source_rows'].append({'source_row':rowdata['source_row'], 'text':rowdata['text'],
                                                'type':'metadata', 'metadata':dict(entry['metadata'])})
                sheet.append(['Tổng cộng'])
                for field,label in METADATA_LABELS.items():
                    if entry['metadata'].get(field):sheet.append([label+': '+entry['metadata'][field]])
                sheet.append(header)
                continue
            cells=entry['cells']
            if not cells['name'] and not any(cells[field] for field in ('quantity','price','amount')):continue
            rawcells=dict(cells)
            if not cells['name']:
                cells={**cells,'name':'[OCR: thiếu tên hàng]',
                       'quantity':'[OCR: thiếu tên hàng] '+str(cells['quantity'])}
                _issue(docreport,'error',f'Dòng {sheet.max_row+1}: ảnh có số liệu nhưng thiếu tên hàng.',row=sheet.max_row+1,field='name',source_row=rowdata['source_row'])
            for field in ('quantity','price','amount'):
                if str(cells.get(field,'')).startswith('='):cells={**cells,field:'[OCR: số không hợp lệ] '+cells[field]}
            if not any(cells[field] for field in ('quantity','price','amount')):
                cells={**cells,'quantity':'[OCR: thiếu số liệu hoặc dòng chưa rõ]'}
            sheet.append([cells.get(field,'') for field in FIELDS])
            outputrow=sheet.max_row;goods_count+=1
            for cell in sheet[outputrow]:
                if isinstance(cell.value,str):cell.data_type='s'
            source={'row':outputrow,'source_row':rowdata['source_row'],'text':rowdata['text'],'confidence':rowdata['confidence'],
                    'cells':rawcells,'cell_confidence':entry.get('cell_confidence',{}),'bbox':rowdata.get('bbox'),'type':'goods'}
            if entry.get('continuation_rows'):source['continuation_rows']=entry['continuation_rows']
            docreport['source_rows'].append(source)
            if not cells['name']:_issue(docreport,'error',f'Dòng {outputrow}: ảnh có số liệu nhưng thiếu tên hàng.',row=outputrow,field='name',source_row=rowdata['source_row'])
            if number(cells['quantity']) is None:_issue(docreport,'error',f'Dòng {outputrow}: số lượng từ ảnh bị thiếu hoặc chưa đọc được; cần sửa bảng chữ.',row=outputrow,field='quantity',source_row=rowdata['source_row'])
            price,amount=number(cells['price']),number(cells['amount'])
            if price is None and amount is None:_issue(docreport,'error',f'Dòng {outputrow}: chưa đọc được đơn giá hoặc thành tiền; cần sửa bảng chữ.',row=outputrow,field='price',source_row=rowdata['source_row'])
            for field,confidence in entry.get('cell_confidence',{}).items():
                if confidence<65 and cells.get(field):_issue(docreport,'warning',f'Dòng {outputrow}: OCR đọc cột {LABELS.get(field,field)} với độ tin cậy {confidence:g}/100; hãy đối chiếu ảnh.',row=outputrow,field=field,source_row=rowdata['source_row'])
        if not goods_count:
            # A sentinel incomplete row remains an engine error during stateless reanalysis.
            sheet.append(['[OCR: chưa nhận diện được dòng hàng]', '', '', '[OCR: chưa đọc được số lượng]', '', '', ''])
            _issue(docreport,'error','Không có dòng hàng nhận diện được. Hãy sửa bảng chữ hoặc dùng ảnh rõ hơn.')
        if not metadata.get('customer'):_issue(docreport,'warning','Chưa tìm được tên khách hàng trong ảnh; cần nhập và xác nhận.',field='customer')
        if not metadata.get('date'):_issue(docreport,'warning','Chưa tìm được ngày xuất trong ảnh; cần nhập và xác nhận.',field='date')
        sheet.freeze_panes=f'A{header_row+1}'
        for col,width in (('A',38),('B',18),('C',12),('D',18),('E',18),('F',20),('G',30)):sheet.column_dimensions[col].width=width
        docreport['source_rows'].sort(key=lambda row:row['source_row'])
        group_info=customer_info[index-1]
        if group_info:
            docreport.update({key:value for key,value in group_info.items() if key!='sources'})
            docreport['confidence_origin']='edited-text' if group_info['edited_text'] else 'grid-cell-ocr'
            if group_info.get('price_inferred'):_issue(docreport,'warning','Tiêu đề đơn giá chưa đọc rõ; cột đơn giá chung được đề xuất theo bố cục bảng. Hãy đối chiếu ảnh và sửa bảng chữ nếu cột chưa đúng.',field='price')
            if group_info.get('header_inferred'):_issue(docreport,'warning','Một số tiêu đề Tổng chưa đọc rõ; nhóm khách hàng được đề xuất theo các cột lặp lại. Hãy đối chiếu tên khách hàng, số lượng và thành tiền với ảnh.',field='quantity')
            if not group_info['edited_text']:
                for row,source in zip([row for row in docreport['source_rows'] if row.get('type')=='goods'],group_info['sources']):
                    row['confidence']=source['confidence']
                    if source['confidence']<65:_issue(docreport,'warning',f"Dòng {row['row']}: ô trong ảnh chưa đọc chắc, hãy đối chiếu ảnh.",row=row['row'])
        report['documents'].append(docreport);report['image_info'][title]=docreport
    output=io.BytesIO();wb.save(output)
    return output.getvalue(),report
